#!/usr/bin/env python3
"""Uji menu **SPJ TANGCER** — dari daftar penerima jadi berkas siap tanda tangan.

Yang dibuktikan uji ini:

1. isian SPJ punya nilai bawaan yang benar (biaya personal 100.000, volume 6 bulan, jadi
   600.000 per siswa) dan semester «Ganjil/1» diseragamkan;
2. baris SPJ diambil dari data penerima yang **tersimpan** — nomor rekening, kelas, jenis
   kelamin, dan nama orang tua ikut terbawa;
3. enam berkas SPJ dibuat sungguhan: dua DOCX, satu XLSX, satu PDF, dan dua lembar
   dokumentasi — nama berkasnya sama seperti berkas kerja sekolah;
4. berkas DOCX hasilnya **sah** (zip + XML-nya bisa dibuka ulang, dan kalau pustaka
   pihak ketiga ``python-docx`` terpasang, isinya dibaca lewat pustaka itu juga);
5. XLSX hasil punya 11 kolom seperti berkas «DAFTAR PENERIMA.xlsx» dan kolom JUMLAH
   memakai rumus ``=biaya×volume``;
6. PDF hasil bisa dibaca kembali oleh pembaca PDF aplikasi;
7. unduhan hanya melayani berkas yang berada di folder paketnya (menolak ``../``);
8. halaman ``/tangcer/spj`` bisa dibuka, menyusun berkas lewat formulir, mengunduh, dan
   menghapus paketnya — semuanya lewat HTTP sebagai petugas.

Jalankan:

    .venv/bin/python scripts/uji_spj.py
    .venv/bin/python scripts/uji_spj.py --berkas-ke pratinjau/ronde60/berkas
"""

from __future__ import annotations

import argparse
import os
import shutil
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

AKAR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(AKAR))

os.environ.setdefault("SM_DB", str(AKAR / "data" / "sm.sqlite3"))

from app import pdf_baca, spj, tangcer  # noqa: E402

TAHUN_UJI = "2099"
SEKOLAH = "SMP NEGERI 2 TANGERANG"
KEPALA = "Drs. Uji Kepala Sekolah, M.Pd."
NIP = "19700101 199001 1 001"
FOTO = "FOTOCOPY BUKU TABUNGAN dan Dokumentasi Tanda Terima ( SAMPEL 10 ANAK ).docx"

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

HASIL: list[tuple[str, bool, str]] = []
SKENARIO: list[str] = []
_SKENARIO_KINI = ""


def skenario(nama: str) -> None:
    global _SKENARIO_KINI
    _SKENARIO_KINI = nama
    SKENARIO.append(nama)
    print(f"\n=== {nama} ===", flush=True)


def cek(nama: str, syarat: bool, keterangan: str = "") -> None:
    HASIL.append((f"{_SKENARIO_KINI} → {nama}", bool(syarat), keterangan))
    tanda = "  OK  " if syarat else " GAGAL"
    print(f"[{tanda}] {nama}" + ("" if syarat or not keterangan else f" — {keterangan}"),
          flush=True)


# --------------------------------------------------------------------------- #
# Data penerima contoh
# --------------------------------------------------------------------------- #
NAMA_CONTOH = [
    ("0091234501", "Achmad Fauzi Ramadhan", "7A", "L", "SITI AISAH", "0148466191100"),
    ("0091234502", "Bella Safitri", "7B", "P", "ARIE KUSUMAWATI", "0153300097100"),
    ("0091234503", "Candra Wijaya Kusuma", "8A", "L", "SITI NURHALIMAH", "0115300636100"),
    ("0091234504", "Dewi Anggraeni", "9C", "P", "RUKMINI", "0160602716100"),
    ("0091234505", "Eko Prasetyo", "7A", "L", "SUPARMIN", "0144810015100"),
    ("0091234506", "Fitriani Lestari", "8B", "P", "YULIANAH", "0160602767100"),
    ("0091234507", "Gilang Ramadhan", "7C", "L", "MULYADI", "0082939970100"),
    ("0091234508", "Hana Salsabila", "9A", "P", "NOVIANI WIJAYA", "0121327831100"),
    ("0091234509", "Irfan Maulana", "8C", "L", "HADI BUDOYO", "0153300143100"),
    ("0091234510", "Jihan Aulia", "7D", "P", "SOPIAN", "0160602694100"),
    ("0091234511", "Kevin Ardiansyah", "9B", "L", "DODI SOPANDI", "0160602686100"),
    ("0091234512", "Laila Nur Fadilah", "8D", "P", "BAHRUDIN", "0153300208100"),
]


def siapkan_penerima() -> int:
    """Simpan 12 penerima contoh untuk tahun uji (dibersihkan lagi di akhir)."""
    item = []
    for nisn, nama, rombel, kelamin, orangtua, rekening in NAMA_CONTOH:
        item.append(tangcer.BarisPenerima(
            nama=nama, nisn=nisn, nik=f"3671{nisn[4:]}0001", rombel=rombel, sekolah=SEKOLAH,
            rekening=rekening, nama_ayah=orangtua, kelamin=kelamin, tempat_lahir="TANGERANG",
            tanggal_lahir="01-01-2011", jenis="TANGCER", nominal=600_000))
    baru, _ = tangcer.simpan_banyak(0, item, tahun=TAHUN_UJI, semester="1", tahap="1")
    return baru


def bersihkan() -> None:
    tangcer.hapus_penerima(TAHUN_UJI, "1")
    for paket in spj.paket_tersimpan(batas=50):
        if paket["token"].startswith("spj-"):
            spj.hapus_paket(paket["token"])


# --------------------------------------------------------------------------- #
# Pembaca berkas hasil
# --------------------------------------------------------------------------- #
def teks_docx(path: Path) -> list[str]:
    """Semua paragraf + baris tabel sebuah DOCX (juga memastikan zip & XML-nya sah)."""
    with zipfile.ZipFile(path) as berkas:
        rusak = berkas.testzip()
        if rusak:
            raise ValueError(f"berkas rusak: {rusak}")
        akar = ET.fromstring(berkas.read("word/document.xml"))
    keluaran: list[str] = []
    for paragraf in akar.iter(f"{W}p"):
        keluaran.append("".join(t.text or "" for t in paragraf.iter(f"{W}t")))
    for tabel in akar.iter(f"{W}tbl"):
        for baris in tabel.iter(f"{W}tr"):
            keluaran.append(" | ".join("".join(t.text or "" for t in sel.iter(f"{W}t"))
                                       for sel in baris.iter(f"{W}tc")))
    return keluaran


def isi_docx_mentah(path: Path) -> str:
    with zipfile.ZipFile(path) as berkas:
        return berkas.read("word/document.xml").decode("utf-8")


# --------------------------------------------------------------------------- #
# Skenario
# --------------------------------------------------------------------------- #
def uji_isian() -> None:
    skenario("1. Isian SPJ & angka bawaan")
    data = spj.data_dari({"tahun": "2026", "semester": "1"})
    cek("biaya personal bawaan 100.000", data.biaya_personal == 100_000,
        str(data.biaya_personal))
    cek("volume bawaan 6 bulan (satu semester)", data.volume == 6, str(data.volume))
    cek("periode semester 1 = Januari s/d Juni", data.periode == "Januari s/d Juni",
        data.periode)
    cek("semester «Ganjil» dianggap semester 1", spj.samakan_semester("Ganjil") == "1",
        spj.samakan_semester("Ganjil"))
    cek("semester «Genap/2» dianggap semester 2", spj.samakan_semester("Genap") == "2",
        spj.samakan_semester("Genap"))
    cek("semester 2 memakai periode Juli s/d Desember",
        spj.data_dari({"semester": "2"}).periode == "Juli s/d Desember")
    cek("biaya & volume yang aneh dikoreksi ke nilai bawaan",
        spj.data_dari({"biaya_personal": "0", "volume": "abc"}).biaya_personal == 100_000,
        "isian kosong/huruf harus jatuh ke bawaan")
    cek("tahun diisi dari tahun ajaran sekolah bila kosong",
        bool(spj.data_dari({}).tahun), spj.data_dari({}).tahun)
    cek("biaya personal teks «Rp 100.000» tetap terbaca 100000",
        spj.data_dari({"biaya_personal": "Rp 100.000"}).biaya_personal == 100_000)
    cek("tanggal tanda tangan menyebut bulan Indonesia (mis. «Oktober»)",
        any(bulan in spj.data_dari({"kota": "Tangerang"}).tanggal_teks
            for bulan in ("Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
                          "Agustus", "September", "Oktober", "November", "Desember")),
        spj.data_dari({"kota": "Tangerang"}).tanggal_teks)


def uji_penerima() -> None:
    skenario("2. Baris SPJ dari data penerima tersimpan")
    baru = siapkan_penerima()
    cek("12 penerima contoh tersimpan", baru == 12, f"tersimpan {baru}")
    data = spj.data_dari({"tahun": TAHUN_UJI, "semester": "1", "kepala": KEPALA, "nip": NIP})
    baris = spj.baris_spj(data)
    cek("baris SPJ diambil dari data tersimpan", len(baris) == 12, f"{len(baris)} baris")
    cek("catatan menyebut dari mana datanya diambil",
        any("data tersimpan" in satu for satu in data.catatan), str(data.catatan))
    pertama = baris[0]
    cek("nomor rekening terbawa", pertama["rekening"] == "0148466191100",
        pertama["rekening"])
    cek("kelas/rombel terbawa", pertama["rombel"] == "7A", pertama["rombel"])
    cek("jenis kelamin terbawa", pertama["kelamin"] == "L", pertama["kelamin"])
    cek("nama orang tua terbawa", pertama["orangtua"] == "SITI AISAH", pertama["orangtua"])
    cek("jumlah per siswa = biaya × volume = 600.000", pertama["jumlah"] == 600_000,
        str(pertama["jumlah"]))
    ringkas = spj.ringkas_spj(data, baris)
    cek("ringkasan: 12 siswa, total 7.200.000",
        ringkas["jumlah_siswa"] == 12 and ringkas["total"] == 7_200_000,
        str(ringkas))
    cek("ringkasan menghitung laki-laki & perempuan",
        ringkas["laki_laki"] + ringkas["perempuan"] == 12,
        f"{ringkas['laki_laki']} L / {ringkas['perempuan']} P")


def uji_paket() -> dict:
    skenario("3. Paket enam berkas SPJ")
    data = spj.data_dari({"tahun": TAHUN_UJI, "semester": "1", "tahap": "1", "kepala": KEPALA,
                          "nip": NIP, "dana_masuk": "1"})
    baris = spj.baris_spj(data)
    paket = spj.buat_paket(data, baris)
    nama_berkas = [satu["berkas"] for satu in paket["berkas"]]
    cek("enam berkas dibuat", len(nama_berkas) == 6, str(nama_berkas))
    cek("nama berkas sama seperti berkas kerja sekolah",
        nama_berkas[0] == "PAKTA INTEGRITAS TANGERANG CERDAS.docx"
        and nama_berkas[1] == "SURAT PERNYATAAN TRANSFER.docx"
        and nama_berkas[2] == "DAFTAR PENERIMA.xlsx"
        and nama_berkas[5] == "DAFTAR SAMPEL DOKUMENTASI.docx", str(nama_berkas))
    cek("lembar fotokopi memakai nama berkas contoh sekolah (dua spasi)",
        FOTO in nama_berkas, str(nama_berkas))
    cek("semua berkas berisi (bukan kosong)",
        all(satu["ukuran"] > 1500 for satu in paket["berkas"]),
        str([(satu["berkas"], satu["ukuran"]) for satu in paket["berkas"]]))
    folder = Path(paket["folder"])
    cek("berkas benar-benar ada di folder paket",
        all((folder / nama).is_file() for nama in nama_berkas), str(folder))
    cek("folder paket memakai token bertanggal", paket["token"].startswith("spj-"),
        paket["token"])
    return paket


def uji_docx_pakta(paket: dict) -> None:
    skenario("4. Pakta Integritas & Surat Pernyataan (DOCX)")
    folder = Path(paket["folder"])
    pakta = folder / "PAKTA INTEGRITAS TANGERANG CERDAS.docx"
    isi = teks_docx(pakta)
    gabung = " ".join(isi)
    cek("berkas DOCX sah dibuka (zip + XML)", "PAKTA INTEGRITAS" in gabung)
    cek("judul dokumen benar", isi and isi[0].strip() == "PAKTA INTEGRITAS", isi[:1])
    cek("nama kepala sekolah terisi", KEPALA in gabung, "nama kepala tidak ada")
    cek("NIP terisi", NIP in gabung, "NIP tidak ada")
    cek("nama sekolah terisi", SEKOLAH in gabung)
    cek("jumlah siswa terisi 12", "sejumlah 12 Siswa" in gabung, gabung[:200])
    cek("tahun & semester disebut", f"Tahun {TAHUN_UJI}" in gabung and "Semester 1" in gabung)
    cek("tempat meterai 10.000 disiapkan", "Materai 10.000" in gabung)
    cek("tidak ada catatan templat yang harus dihapus manual",
        "HAPUS CATATAN INI" not in gabung, "catatan templat ikut terbawa")

    surat = folder / "SURAT PERNYATAAN TRANSFER.docx"
    isi_surat = teks_docx(surat)
    gabung_surat = " ".join(isi_surat)
    cek("surat pernyataan memuat kop sekolah", SEKOLAH in gabung_surat[:200],
        gabung_surat[:120])
    cek("surat menyebut personal 100.000/peserta didik/bulan",
        "Personal : 100.000/peserta didik/bulan" in gabung_surat, "kalimat juknis hilang")
    cek("baris tabel penerimaan & penyaluran terisi total 7.200.000",
        gabung_surat.count("7.200.000") >= 2, gabung_surat[-400:])
    cek("keterangan besaran & semester tertulis",
        f"Semester 1 (Periode Januari s/d Juni) {TAHUN_UJI}" in gabung_surat)
    cek("tabel pernyataan punya 5 kolom",
        any(baris.count(" | ") == 4 for baris in isi_surat),
        "tabel 5 kolom tidak ditemukan")

    # Kalau pustaka pihak ketiga ada, buka lewat pustaka itu — bukti tambahan.
    try:
        import docx  # noqa: PLC0415
        dokumen = docx.Document(str(surat))
        cek("berkas DOCX juga terbaca pustaka python-docx (pihak ketiga)",
            len(dokumen.tables) >= 2 and dokumen.tables[0].cell(1, 3).text == "7.200.000",
            f"{len(dokumen.tables)} tabel")
    except ImportError:
        cek("(pustaka python-docx tidak terpasang — pemeriksaan pihak ketiga dilewati)", True)


def uji_xlsx(paket: dict) -> None:
    skenario("5. Daftar Penerima (XLSX) — bentuk berkas sekolah")
    import openpyxl  # noqa: PLC0415

    lembar = openpyxl.load_workbook(Path(paket["folder"]) / "DAFTAR PENERIMA.xlsx").active
    kepala = [lembar.cell(row=4, column=kolom).value for kolom in range(1, 12)]
    cek("judul lembar menyebut TANGERANG CERDAS dan jenjang SMP",
        "TANGERANG CERDAS" in str(lembar["A1"].value) and "SMP" in str(lembar["A1"].value))
    cek("baris kedua menyebut semester & tahun anggaran",
        f"SEMESTER 1" in str(lembar["A2"].value).upper() and TAHUN_UJI in str(lembar["A2"].value),
        str(lembar["A2"].value))
    cek("11 kolom persis seperti berkas «DAFTAR PENERIMA.xlsx»",
        kepala == ["NO", "NOMOR REKENING", "NAMA", "KELAS", "L/P", "NAMA ORANGTUA",
                   "NAMA SEKOLAH", "BIAYA PERSONAL", "VOLUME", "JUMLAH", "TANDA TANGAN"],
        str(kepala))
    cek("baris nomor kolom «10 = 8x9» ada", lembar["J5"].value == "10 = 8x9",
        str(lembar["J5"].value))
    cek("kolom JUMLAH memakai rumus =H×I", str(lembar["J6"].value) == "=H6*I6",
        str(lembar["J6"].value))
    cek("biaya personal 100000 & volume 6 terisi",
        lembar["H6"].value == 100000 and lembar["I6"].value == 6,
        f"{lembar['H6'].value} × {lembar['I6'].value}")
    cek("nama & rekening penerima pertama benar",
        lembar["C6"].value == "Achmad Fauzi Ramadhan"
        and lembar["B6"].value == "0148466191100",
        f"{lembar['C6'].value} / {lembar['B6'].value}")
    cek("baris data berjumlah 12 (sesuai penerima tersimpan)",
        lembar["A17"].value == 12 and lembar["A18"].value is None,
        f"baris 17 nomor {lembar['A17'].value}")
    cek("kaki lembar menyiapkan tanda tangan kepala sekolah",
        "Kepala Sekolah" in " ".join(str(lembar.cell(row=r, column=9).value) for r in range(6, 30)),
        "«Kepala Sekolah» tidak ditemukan")


def uji_pdf_dan_dokumentasi(paket: dict) -> None:
    skenario("6. Daftar Penerima (PDF) & lembar dokumentasi")
    folder = Path(paket["folder"])
    isi = (folder / "DAFTAR PENERIMA.pdf").read_bytes()
    hasil = pdf_baca.baca_pdf(isi)
    teks = " ".join(t for _, _, t, _ in hasil.halaman[0])
    cek("PDF bisa dibaca pembaca bawaan aplikasi", hasil.jumlah_halaman >= 1,
        f"{hasil.jumlah_halaman} halaman")
    cek("PDF berkepala (nama sekolah & semester)", SEKOLAH in teks and "Semester 1" in teks,
        teks[:120])
    cek("PDF memuat nama penerima pertama", "Achmad Fauzi Ramadhan" in teks)
    cek("PDF memuat nomor rekening", "0148466191100" in teks)
    cek("PDF memuat total 7.200.000", "7.200.000" in teks, teks[-160:])

    sampel = teks_docx(folder / "DAFTAR SAMPEL DOKUMENTASI.docx")
    gabung = " ".join(sampel)
    cek("daftar sampel memuat 10 anak (sampel bawaan)",
        sum(1 for satu in sampel if satu.startswith("1 | ") or satu.split(" | ")[0].isdigit()
            and "|" in satu) >= 10, "baris sampel kurang")
    cek("daftar sampel memuat nama & rekening penerima",
        "Achmad Fauzi Ramadhan" in gabung and "0148466191100" in gabung)
    cek("daftar sampel memuat kolom centang", "☐" in gabung)
    cek("daftar sampel menyebut jumlah penerima keseluruhan",
        "12 siswa" in gabung, "jumlah penerima tidak disebut")

    fotokopi = teks_docx(folder / FOTO)
    gabung_foto = " ".join(fotokopi)
    cek("lembar fotokopi memakai dua judul seperti berkas sekolah",
        "FOTOCOPY IDENTITAS BUKU TABUNGAN" in gabung_foto
        and "FOTO DOKUMENTASI SAAT TANDA TANGAN" in gabung_foto)


def uji_unduhan_aman(paket: dict) -> None:
    skenario("7. Unduhan hanya dari folder paketnya")
    token = paket["token"]
    nama = "PAKTA INTEGRITAS TANGERANG CERDAS.docx"
    cek("berkas paket bisa ditemukan lewat token & namanya",
        spj.berkas_spj(token, nama) is not None)
    cek("nama berkas dengan «../» ditolak", spj.berkas_spj(token, "../data/sm.sqlite3") is None)
    cek("nama berkas bersubfolder ditolak", spj.berkas_spj(token, "sub/berkas.docx") is None)
    cek("token bersubfolder ditolak", spj.berkas_spj("../spj", nama) is None)
    cek("berkas yang tidak ada ditolak", spj.berkas_spj(token, "tidak-ada.docx") is None)


def uji_web(paket: dict) -> None:
    skenario("8. Halaman SPJ lewat HTTP (petugas)")
    import asyncio

    import httpx

    from app.main import app

    token = paket["token"]

    async def alur() -> None:
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://uji",
                                     follow_redirects=True) as klien:
            masuk = await klien.post("/login", data={"mode": "staff", "username": "admin",
                                                     "password": "admin123"})
            cek("Masuk sebagai admin berhasil", "Keluar" in masuk.text)

            halaman = await klien.get("/tangcer")
            cek("Halaman TANGCER menyediakan jalan ke SPJ",
                'href="/tangcer/spj"' in halaman.text and "Susun berkas SPJ" in halaman.text)

            spj_hal = await klien.get(f"/tangcer/spj?tahun={TAHUN_UJI}&semester=1")
            cek("Halaman SPJ terbuka", spj_hal.status_code == 200
                and "SPJ Tangerang Cerdas" in spj_hal.text, f"status {spj_hal.status_code}")
            cek("Pratinjau menampilkan penerima dari data tersimpan",
                "Achmad Fauzi Ramadhan" in spj_hal.text, "nama penerima tidak muncul")
            cek("Pratinjau menampilkan rekening & nama orang tua",
                "0148466191100" in spj_hal.text and "SITI AISAH" in spj_hal.text)
            cek("Halaman menyebut total Rp 7.200.000", "Rp 7.200.000" in spj_hal.text)
            cek("Formulir isian SPJ lengkap",
                all(f'name="{nama}"' in spj_hal.text for nama in
                    ("tahun", "semester", "kepala", "nip", "biaya_personal", "volume",
                     "sampel", "dana_masuk")))

            buat = await klien.post("/tangcer/spj", data={
                "tahun": TAHUN_UJI, "semester": "1", "tahap": "1", "kepala": KEPALA,
                "nip": NIP, "jabatan": "Kepala Sekolah", "dana_masuk": "1",
                "biaya_personal": "100000", "volume": "6", "sampel": "10"})
            cek("Formulir menyusun berkas & membuka halaman hasil",
                buat.status_code == 200 and "Berkas SPJ siap unduh" in buat.text,
                f"status {buat.status_code}")
            cek("Pesan sukses menyebut jumlah berkas & siswa",
                "berkas untuk 12 siswa" in buat.text, "pesan sukses tidak sesuai")

            unduh = await klien.get("/tangcer/spj/unduh", params={
                "token": token, "berkas": "PAKTA INTEGRITAS TANGERANG CERDAS.docx"})
            cek("Berkas DOCX bisa diunduh dari halaman",
                unduh.status_code == 200
                and "wordprocessingml" in unduh.headers.get("content-type", ""),
                f"status {unduh.status_code} · {unduh.headers.get('content-type')}")
            cek("Unduhan DOCX isinya sah (zip terbaca)",
                unduh.content[:2] == b"PK" and len(unduh.content) > 1500,
                f"{len(unduh.content)} bita")

            xlsx = await klien.get("/tangcer/spj/unduh", params={
                "token": token, "berkas": "DAFTAR PENERIMA.xlsx"})
            cek("Berkas XLSX bisa diunduh dengan media type Excel",
                xlsx.status_code == 200 and "spreadsheetml" in xlsx.headers.get("content-type", ""))

            bukan = await klien.get("/tangcer/spj/unduh", params={
                "token": "spj-tidak-ada", "berkas": "apa-saja.docx"})
            cek("Unduhan token salah ditolak dengan pesan",
                "tidak ditemukan" in bukan.text.lower(), bukan.text[:120])

            hapus = await klien.post("/tangcer/spj/hapus", data={"token": token})
            cek("Paket SPJ bisa dihapus dari halaman",
                hapus.status_code == 200 and "Paket berkas SPJ dihapus" in hapus.text,
                f"status {hapus.status_code}")
            cek("Berkas paket benar-benar hilang setelah dihapus",
                spj.berkas_spj(token, "DAFTAR PENERIMA.pdf") is None)

    asyncio.run(alur())


def simpan_contoh(paket: dict, tujuan: Path) -> None:
    if tujuan.exists():
        shutil.rmtree(tujuan)
    shutil.copytree(paket["folder"], tujuan)
    print(f"\nContoh berkas SPJ disimpan di {tujuan}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Uji menu SPJ TANGCER")
    parser.add_argument("--berkas-ke", type=Path, default=None,
                        help="salin berkas hasil ke folder ini (untuk bukti)")
    args = parser.parse_args()

    tangcer.hapus_penerima(TAHUN_UJI, "1")
    try:
        uji_isian()
        uji_penerima()
        paket = uji_paket()
        uji_docx_pakta(paket)
        uji_xlsx(paket)
        uji_pdf_dan_dokumentasi(paket)
        uji_unduhan_aman(paket)
        if args.berkas_ke:
            simpan_contoh(paket, args.berkas_ke)
        uji_web(paket)
    finally:
        bersihkan()

    gagal = [satu for satu in HASIL if not satu[1]]
    print(f"\n[SELESAI] {len(HASIL) - len(gagal)} pemeriksaan lolos pada {len(SKENARIO)} skenario")
    if gagal:
        print(f"[GAGAL] {len(gagal)} pemeriksaan tidak lolos:")
        for nama, _, keterangan in gagal:
            print(f"   - {nama} ({keterangan})")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
