#!/usr/bin/env python3
"""Uji menyeluruh menu **Guru & Tendik (GTK)** — impor berkas mentah Dapodik.

Berkas «daftar guru» / «daftar tendik» dari Dapodik langsung dibaca apa adanya
(kop baris 1-4 dilewati, judul kolom ditemukan sendiri), dipratinjau, lalu
disimpan sebagai daftar pegawai. Uji ini membuktikan bahwa:

1. judul kolom ditemukan walau kopnya bergeser/tambah-baris;
2. seluruh 50 kolom Dapodik terpetakan dan nilainya apa adanya (tanggal → ISO,
   Ya/Tidak dinormalisasi, NPWP dipangkas);
3. baris bermasalah dilaporkan (nama kosong = galat, tanpa NUPTK/NIP/NIK =
   peringatan, JK di luar L/P = peringatan);
4. penyimpanan tidak menggandakan baris (kunci NUPTK → NIP → NIK), unggahan
   ulang hanya memperbarui yang berubah;
5. pencarian, filter jenis/status, dan ekspor CSV/XLSX bekerja;
6. berkas asli sekolah (daftar guru 44 baris + daftar tendik 13 baris) terbaca
   penuh dan tersimpan sebagai 57 pegawai.

Jalankan::

    .venv/bin/python scripts/uji_gtk.py
"""

from __future__ import annotations

import asyncio
import os
import sys
import tempfile
import traceback
from pathlib import Path

AKAR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(AKAR))

# --- Database sementara: HARUS diatur sebelum mengimpor modul app ---
_SEMENTARA = Path(tempfile.mkdtemp(prefix="sm-uji-gtk-"))
os.environ["SM_DATA_DIR"] = str(_SEMENTARA / "data")
os.environ["SM_DB_PATH"] = str(_SEMENTARA / "uji.sqlite3")
os.environ["SM_AUTO_SEED"] = "0"

from app import db, gtk, migrations  # noqa: E402

migrations.init_database()

HASIL: list[tuple[str, bool, str]] = []
SKENARIO: list[str] = []
_SKENARIO_KINI = ""


def skenario(nama: str) -> None:
    global _SKENARIO_KINI
    _SKENARIO_KINI = nama
    SKENARIO.append(nama)
    print(f"\n=== {nama} ===", flush=True)


def cek(nama: str, syarat: bool, keterangan: str = "") -> None:
    HASIL.append((f"{_SKENARIO_KINI} → {nama}", bool(syarat), keterangan))
    tanda = "  OK  " if syarat else " GAGAL"
    print(f"[{tanda}] {nama}" + ("" if syarat or not keterangan else f" — {keterangan}"),
          flush=True)


KEPALA_GTK = [spec.label for spec in gtk.GTK_FIELDS]

KOP = [
    ["Daftar Guru"],
    ["SMP NEGERI 2 TANGERANG"],
    ["Kecamatan Kec. Tangerang, Kabupaten Kota Tangerang, Provinsi Prov. Banten"],
    ["Tanggal Unduh: 2026-10-22 10:05:18", None,
     "Pengunduh: Operator (smpntangerang2@gmail.com)"],
]


def _baris_contoh(**ubah) -> list:
    dasar = {
        "Nama": "Siti Rahmawati", "NUPTK": "1234567890123456", "JK": "P",
        "Tempat Lahir": "Tangerang", "Tanggal Lahir": "1985-03-12",
        "NIP": "198503122010012001", "Status Kepegawaian": "PNS",
        "Jenis PTK": "Guru", "Agama": "Islam",
        "Alamat Jalan": "Jl. Merdeka No. 10", "RT": "1", "RW": "2",
        "Desa/Kelurahan": "Suka Maju", "Kecamatan": "Kec. Tangerang",
        "Kode Pos": "15111", "HP": "081234567890", "Email": "siti@example.com",
        "Tugas Tambahan": "Wali Kelas", "TMT Pengangkatan": "2010-01-01",
        "Pangkat Golongan": "III/b", "Sumber Gaji": "APBN",
        "Nama Ibu Kandung": "Ibu Siti", "Status Perkawinan": "Kawin",
        "NPWP": "  123456789012345  ", "NIK": "3671000000000001",
        "No KK": "3671000000000002",
    }
    dasar.update(ubah)
    return [dasar.get(label) for label in KEPALA_GTK]


def buat_xlsx_contoh(path: Path, kop: list | None = None,
                     baris_data: list | None = None) -> Path:
    import openpyxl

    buku = openpyxl.Workbook()
    sheet = buku.active
    sheet.title = "Daftar guru"
    for baris in (KOP if kop is None else kop):
        sheet.append(baris)
    sheet.append(KEPALA_GTK)
    for baris in (baris_data if baris_data is not None else [_baris_contoh()]):
        sheet.append(baris)
    buku.save(path)
    return path


# --------------------------------------------------------------------------- #
# 1. Parser & temuannya
# --------------------------------------------------------------------------- #
skenario("1. Parser menemukan judul & memetakan kolom")
contoh = buat_xlsx_contoh(_SEMENTARA / "contoh-gtk.xlsx")
sheet, hasil, label = gtk.baca_berkas(contoh)
cek("judul kolom ditemukan pada baris 5", hasil.header_row == 5, f"dapat {hasil.header_row}")
cek("data mulai baris 6", hasil.data_start_row == 6)
cek("jenis dikenali dari judul", hasil.kind == "daftar_gtk", hasil.kind)
cek("berkas dikenali sebagai daftar guru", hasil.jenis_berkas == "guru", hasil.jenis_berkas)
cek("50 kolom terpetakan", hasil.mapped_field_count == 50, f"dapat {hasil.mapped_field_count}")
cek("sekolah terbaca dari kop", hasil.metadata.get("sekolah") == "SMP NEGERI 2 TANGERANG",
    repr(hasil.metadata.get("sekolah")))
cek("satu baris data terbaca", len(hasil.records) == 1, f"dapat {len(hasil.records)}")
nilai = hasil.records[0].values
cek("nama dibaca apa adanya (kapital)", nilai.get("nama") == "SITI RAHMAWATI", repr(nilai.get("nama")))
cek("tanggal lahir → ISO", nilai.get("tanggal_lahir") == "1985-03-12",
    repr(nilai.get("tanggal_lahir")))
cek("NPWP dipangkas", nilai.get("npwp") == "123456789012345", repr(nilai.get("npwp")))
cek("tanpa temuan pada baris bersih", not hasil.records[0].issues,
    repr([t.message for t in hasil.records[0].issues]))

skenario("2. Judul tetap ketemu walau kop bergeser")
kop_geser = [["LAPORAN DAPODIK"], [], *KOP, ["Catatan: data per semester ganjil"]]
geser = buat_xlsx_contoh(_SEMENTARA / "contoh-geser.xlsx", kop=kop_geser)
_, hasil_geser, _ = gtk.baca_berkas(geser)
cek("judul ditemukan pada baris 8", hasil_geser.header_row == 8, f"dapat {hasil_geser.header_row}")
cek("data tetap 1 baris", len(hasil_geser.records) == 1)

skenario("3. Baris bermasalah dilaporkan")
bermasalah = buat_xlsx_contoh(
    _SEMENTARA / "contoh-bermasalah.xlsx",
    baris_data=[
        # Tiap baris memakai NUPTK/NIP/NIK berbeda supaya kunci pengenalnya unik.
        _baris_contoh(**{"Nama": "X", "NUPTK": "111", "NIP": "111", "NIK": "111"}),
        _baris_contoh(**{"Nama": "Tanpa Identitas", "NUPTK": None, "NIP": None,
                         "NIK": None}),                                  # tanpa kunci → warning
        _baris_contoh(**{"Nama": "Kelamin Aneh", "JK": "Laki", "NUPTK": "333",
                         "NIP": "333", "NIK": "333"}),                   # JK → warning
        _baris_contoh(**{"Nama": "Tanggal Aneh", "Tanggal Lahir": "kemarin",
                         "NUPTK": "444", "NIP": "444", "NIK": "444"}),   # tgl → warning
        _baris_contoh(**{"Nama": "Tanpa NUPTK 2", "NUPTK": None, "NIP": "555",
                         "NIK": "3671000000000555"}),                   # info saja
    ],
)
_, hasil_rusak, _ = gtk.baca_berkas(bermasalah)
ringkas = gtk.ringkas_temuan(hasil_rusak.issues)
cek("5 baris tetap terbaca", len(hasil_rusak.records) == 5)
cek("nama pendek = galat", hasil_rusak.records[0].has_error)
cek("baris lain tanpa galat", not any(r.has_error for r in hasil_rusak.records[1:]))
cek("ada peringatan kunci (tanpa NUPTK/NIP/NIK)",
    any("Tanpa NUPTK/NIP/NIK" in t.message for t in hasil_rusak.records[1].issues))
cek("ada peringatan JK", any(t.field == "jk" for t in hasil_rusak.records[2].issues))
cek("ada peringatan tanggal", any(t.field == "tanggal_lahir" for t in hasil_rusak.records[3].issues))
cek("tanpa NUPTK (ada NIP/NIK) hanya info",
    [t.level for t in hasil_rusak.records[4].issues] == ["info"],
    repr([(t.level, t.message) for t in hasil_rusak.records[4].issues]))
cek("ringkas temuan: 1 galat", ringkas["error"] == 1, repr(ringkas))

skenario("4. Kunci pengenal baris")
cek("NUPTK diprioritaskan", gtk.kunci_ptk({"nuptk": " 111 ", "nip": "222", "nik": "333"}) == "111")
cek("NIP bila NUPTK kosong", gtk.kunci_ptk({"nuptk": "", "nip": "222", "nik": "333"}) == "222")
cek("NIK bila keduanya kosong", gtk.kunci_ptk({"nip": None, "nik": "333"}) == "333")
cek("nama+lahir bila semua kosong",
    gtk.kunci_ptk({"nama": "Ani", "tanggal_lahir": "2000-01-01"}) == "nama:ANI|lahir:2000-01-01")

# --------------------------------------------------------------------------- #
# 2. Penyimpanan
# --------------------------------------------------------------------------- #
skenario("5. Simpan: baru, ganda, dan diperbarui")
id_impor = gtk.buat_impor(nama_berkas="contoh-gtk.xlsx", path=contoh, ext=".xlsx",
                          ukuran=contoh.stat().st_size, sheet=sheet, hasil=hasil,
                          fmt=label, actor="uji")
cek("kepala impor tersimpan", gtk.get_impor(id_impor) is not None)
simpan1 = gtk.simpan_banyak(id_impor, hasil.records, source_file="contoh-gtk.xlsx")
cek("1 pegawai baru", simpan1["baru"] == 1, repr(simpan1))
gtk.tandai_impor(id_impor, status="tersimpan", baru=simpan1["baru"],
                 diperbarui=simpan1["diperbarui"], dilewati=simpan1["dilewati"],
                 gagal=simpan1["gagal"])
kepala = gtk.get_impor(id_impor)
cek("impor ditandai tersimpan", kepala["status"] == "tersimpan" and kepala["disimpan_baru"] == 1)
simpan2 = gtk.simpan_banyak(id_impor, hasil.records, source_file="contoh-gtk.xlsx")
cek("simpan ulang: tidak ada yang baru", simpan2["baru"] == 0, repr(simpan2))
cek("simpan ulang: dilewati (sudah sama)", simpan2["dilewati"] == 1, repr(simpan2))
cek("tetap 1 pegawai di basis data",
    db.query_value("SELECT COUNT(*) FROM gtk_pegawai") == 1)

# Ubah HP-nya lalu simpan lagi → diperbarui, bukan digandakan.
hasil.records[0].values["hp"] = "089999999999"
simpan3 = gtk.simpan_banyak(id_impor, hasil.records, source_file="contoh-gtk.xlsx")
cek("nilai berubah → diperbarui", simpan3["diperbarui"] == 1, repr(simpan3))
cek("HP baru tersimpan",
    db.query_value("SELECT hp FROM gtk_pegawai WHERE nuptk = '1234567890123456'") == "089999999999")
cek("tetap 1 pegawai (tidak ganda)",
    db.query_value("SELECT COUNT(*) FROM gtk_pegawai") == 1)

# Baris bergalat tidak ikut tersimpan.
_, hasil_rusak, _ = gtk.baca_berkas(bermasalah)
simpan_rusak = gtk.simpan_banyak(id_impor, hasil_rusak.records, source_file="rusak.xlsx")
cek("baris galat gagal (tidak tersimpan)", simpan_rusak["gagal"] == 1, repr(simpan_rusak))
cek("baris baik tetap tersimpan", simpan_rusak["baru"] == 4, repr(simpan_rusak))

skenario("6. Ringkas, cari, filter, ekspor")
ringkas_db = gtk.ringkas_pegawai()
cek("ringkas total 5", ringkas_db["total"] == 5, repr(ringkas_db["total"]))
cek("ringkas guru 5", ringkas_db["guru"] == 5)
cek("daftar_jenis memuat Guru", gtk.daftar_jenis() == ["Guru"], repr(gtk.daftar_jenis()))
cek("cari nama menyaring",
    len(gtk.daftar_pegawai(cari="rahma")) == 1 and
    gtk.daftar_pegawai(cari="rahma")[0]["nama"] == "SITI RAHMAWATI")
cek("cari NIP menyaring",
    len(gtk.daftar_pegawai(cari="198503122010012001")) == 1)
cek("cari yang tidak ada → kosong", gtk.daftar_pegawai(cari="zzz-tidak-ada") == [])
cek("filter jenis Guru", len(gtk.daftar_pegawai(jenis="Guru")) == 5)
cek("filter jenis lain → kosong",
    gtk.daftar_pegawai(jenis="Tenaga Kependidikan") == [])
csv_isi = gtk.ekspor_csv().decode("utf-8-sig")
cek("CSV memuat judul NUPTK", "NUPTK" in csv_isi.splitlines()[0])
cek("CSV memuat data", "SITI RAHMAWATI" in csv_isi)
cek("CSV 5 baris data", len(csv_isi.splitlines()) == 6, f"dapat {len(csv_isi.splitlines())}")
import io as _io

import openpyxl as _oxl

buku = _oxl.load_workbook(_io.BytesIO(gtk.ekspor_xlsx()), read_only=True)
lembar = buku[buku.sheetnames[0]]
judul = [sel.value for sel in next(lembar.iter_rows(min_row=1, max_row=1))]
cek("XLSX judul memuat NUPTK", "NUPTK" in judul)
cek("XLSX 5 baris data", lembar.max_row == 6, f"dapat {lembar.max_row}")
buku.close()

# --------------------------------------------------------------------------- #
# 3. Berkas asli sekolah
# --------------------------------------------------------------------------- #
skenario("7. Berkas asli: daftar guru 44 + daftar tendik 13")
GURU = AKAR / "daftar-guru-SMP NEGERI 2 TANGERANG-2026-10-22 10_05_18.xlsx"
TENDIK = AKAR / "daftar-tendik-SMP NEGERI 2 TANGERANG-2026-10-22 10_05_34.xlsx"
if not GURU.exists() or not TENDIK.exists():
    cek("berkas asli tersedia", False, "berkas contoh tidak ada di akar repo — dilewati")
else:
    lembar_guru, hasil_guru, _ = gtk.baca_berkas(GURU)
    lembar_tendik, hasil_tendik, _ = gtk.baca_berkas(TENDIK)
    cek("guru: 44 baris", len(hasil_guru.records) == 44, f"dapat {len(hasil_guru.records)}")
    cek("tendik: 13 baris", len(hasil_tendik.records) == 13, f"dapat {len(hasil_tendik.records)}")
    cek("guru: jenis_berkas = guru", hasil_guru.jenis_berkas == "guru")
    cek("tendik: jenis_berkas = tendik", hasil_tendik.jenis_berkas == "tendik")
    cek("guru: 50 kolom dikenali", hasil_guru.mapped_field_count == 50)
    cek("tendik: 50 kolom dikenali", hasil_tendik.mapped_field_count == 50)
    cek("guru: tanpa galat", gtk.ringkas_temuan(hasil_guru.issues)["error"] == 0,
        repr(gtk.ringkas_temuan(hasil_guru.issues)))
    cek("tendik: tanpa galat", gtk.ringkas_temuan(hasil_tendik.issues)["error"] == 0)
    cek("nama guru pertama apa adanya",
        hasil_guru.records[0].values.get("nama") == "ABDUL ROHIM")
    cek("NUPTK guru pertama", hasil_guru.records[0].values.get("nuptk") == "9447760661130142")
    kunci_ganda = len({gtk.kunci_ptk(r.values) for r in hasil_guru.records + hasil_tendik.records})
    cek("57 kunci unik (tidak ada ganda)", kunci_ganda == 57, f"dapat {kunci_ganda}")

    # Salin dulu ke folder sementara: hapus_impor ikut menghapus berkas sumbernya,
    # dan berkas asli milik sekolah tidak boleh ikut terhapus.
    import shutil as _shutil

    guru_salin = _SEMENTARA / GURU.name
    tendik_salin = _SEMENTARA / TENDIK.name
    _shutil.copy(GURU, guru_salin)
    _shutil.copy(TENDIK, tendik_salin)
    id_guru = gtk.buat_impor(nama_berkas=GURU.name, path=guru_salin, ext=".xlsx",
                             ukuran=guru_salin.stat().st_size, sheet=lembar_guru,
                             hasil=hasil_guru, fmt="Excel", actor="uji")
    id_tendik = gtk.buat_impor(nama_berkas=TENDIK.name, path=tendik_salin, ext=".xlsx",
                               ukuran=tendik_salin.stat().st_size, sheet=lembar_tendik,
                               hasil=hasil_tendik, fmt="Excel", actor="uji")
    simpan_guru = gtk.simpan_banyak(id_guru, hasil_guru.records, source_file=GURU.name)
    simpan_tendik = gtk.simpan_banyak(id_tendik, hasil_tendik.records, source_file=TENDIK.name)
    cek("44 guru tersimpan", simpan_guru["baru"] == 44, repr(simpan_guru))
    cek("13 tendik tersimpan", simpan_tendik["baru"] == 13, repr(simpan_tendik))
    cek("total 62 pegawai (5 contoh + 57 asli)",
        db.query_value("SELECT COUNT(*) FROM gtk_pegawai") == 62)
    # Unggah ulang berkas guru → tidak ada yang baru.
    simpan_ulang = gtk.simpan_banyak(id_guru, hasil_guru.records, source_file=GURU.name)
    cek("unggah ulang: 0 baru, 44 sama",
        simpan_ulang["baru"] == 0 and simpan_ulang["dilewati"] == 44, repr(simpan_ulang))

    skenario("8. Hapus impor & hapus pegawai")
    gtk.hapus_impor(id_tendik)
    cek("hapus impor tendik → 49 pegawai (5 contoh + 44 guru)",
        db.query_value("SELECT COUNT(*) FROM gtk_pegawai") == 49)
    cek("kepala impor ikut terhapus", gtk.get_impor(id_tendik) is None)
    cek("berkas asli sekolah tidak ikut terhapus", GURU.exists() and TENDIK.exists())
    jumlah = gtk.hapus_pegawai(jenis="Guru")
    cek("hapus massal jenis Guru → 49 baris", jumlah == 49, f"dapat {jumlah}")
    cek("tidak ada sisa", db.query_value("SELECT COUNT(*) FROM gtk_pegawai") == 0)

# --------------------------------------------------------------------------- #
# 4. Alur HTTP
# --------------------------------------------------------------------------- #
skenario("9. Alur HTTP: unggah → pratinjau → simpan → daftar → detail → hapus")


async def alur_http() -> None:
    import httpx

    from app.main import app

    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://uji",
                                 follow_redirects=True) as klien:
        await klien.post("/login", data={"mode": "staff", "username": "admin",
                                         "password": "admin123"})
        halaman = await klien.get("/gtk")
        cek("GET /gtk 200", halaman.status_code == 200, f"dapat {halaman.status_code}")
        cek("formulir unggah GTK tampil", 'id="berkas-gtk"' in halaman.text)
        cek("menu samping memuat Guru & Tendik", "Guru &" in halaman.text)

        with open(contoh, "rb") as berkas:
            unggah = await klien.post("/gtk/unggah", files={"berkas": ("contoh-gtk.xlsx", berkas.read())})
        cek("unggah → pratinjau", "/gtk/impor/" in str(unggah.url), str(unggah.url))
        cek("tombol Simpan tampil", "Simpan ke daftar pegawai" in unggah.text)
        cek("nama contoh tampil di pratinjau", "SITI RAHMAWATI" in unggah.text)
        iid = str(unggah.url).rstrip("/").rsplit("/", 1)[-1]
        simpan = await klien.post(f"/gtk/impor/{iid}/simpan")
        cek("simpan via HTTP", "pegawai baru disimpan" in simpan.text
            or "sudah sama" in simpan.text, simpan.text[:200])
        daftar = await klien.get("/gtk")
        cek("nama tampil di daftar", "SITI RAHMAWATI" in daftar.text)
        cari = await klien.get("/gtk", params={"cari": "SITI RAHMAWATI"})
        cek("pencarian HTTP menyaring", "SITI RAHMAWATI" in cari.text)
        pid = db.query_value("SELECT id FROM gtk_pegawai WHERE nama = 'SITI RAHMAWATI'")
        detail = await klien.get(f"/gtk/pegawai/{pid}")
        cek("detail 200", detail.status_code == 200)
        cek("detail memuat NUPTK", "1234567890123456" in detail.text)
        cek("detail dikelompokkan", "Kepegawaian" in detail.text and "Pajak &amp; Bank" in detail.text)
        csv_http = await klien.get("/gtk/ekspor.csv")
        cek("ekspor CSV via HTTP", csv_http.status_code == 200 and "NUPTK" in csv_http.text[:600])
        xlsx_http = await klien.get("/gtk/ekspor.xlsx")
        cek("ekspor XLSX via HTTP", xlsx_http.status_code == 200 and len(xlsx_http.content) > 1000)
        hapus = await klien.post(f"/gtk/pegawai/{pid}/hapus")
        cek("hapus satu via HTTP", "dihapus dari daftar pegawai" in hapus.text)
        cek("nama hilang dari daftar", "SITI RAHMAWATI" not in (await klien.get("/gtk")).text)


asyncio.run(alur_http())

# --------------------------------------------------------------------------- #
# Laporan
# --------------------------------------------------------------------------- #
print("\n" + "-" * 78)
gagal = [(nama, ket) for nama, ok, ket in HASIL if not ok]
print(f"[SELESAI] {len(HASIL) - len(gagal)} pemeriksaan lolos pada {len(SKENARIO)} skenario"
      + (f", {len(gagal)} GAGAL" if gagal else ""))
for nama, ket in gagal:
    print(f"  GAGAL: {nama}" + (f" — {ket}" if ket else ""))
if gagal:
    raise SystemExit(1)
