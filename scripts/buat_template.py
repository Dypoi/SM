#!/usr/bin/env python3
"""Membuat berkas contoh (template) impor di folder ``template-import/``.

Template ini memakai susunan kolom yang sama dengan ekspor Dapodik
("Daftar Peserta Didik") namun berisi data fiktif, sehingga aman dibagikan
dan bisa dipakai untuk menguji aplikasi.

Jalankan::

    python scripts/buat_template.py
"""

from __future__ import annotations

import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

from app.dapodik import STUDENT_FIELDS  # noqa: E402

TUJUAN = BASE_DIR / "template-import" / "contoh-template-import.xlsx"

# Data fiktif (bukan data siswa sungguhan)
CONTOH = [
    {
        "nama": "CONTOH SISWA SATU", "nipd": "26260001", "jk": "L", "nisn": "3900000001",
        "tempat_lahir": "TANGERANG", "tanggal_lahir": "2013-01-15",
        "nik": "3671010101130001", "agama": "Islam", "alamat": "Jl. Contoh No. 1",
        "rt": "1", "rw": "2", "dusun": "Contoh", "kelurahan": "Contoh Kelurahan",
        "kecamatan": "Kec. Contoh", "kode_pos": "15111", "jenis_tinggal": "Bersama orang tua",
        "transportasi": "Jalan kaki", "hp": "081200000001", "email": "contoh1@sekolah.sch.id",
        "penerima_kps": "Tidak", "penerima_kip": "Tidak", "layak_pip": "Tidak",
        "ayah_nama": "AYAH CONTOH SATU", "ayah_tahun_lahir": 1980, "ayah_pendidikan": "SMA / sederajat",
        "ayah_pekerjaan": "Karyawan Swasta", "ayah_penghasilan": "Rp. 1,000,000 - Rp. 1,999,999",
        "ayah_nik": "3671010101800001",
        "ibu_nama": "IBU CONTOH SATU", "ibu_tahun_lahir": 1982, "ibu_pendidikan": "SMP / sederajat",
        "ibu_pekerjaan": "Tidak bekerja", "ibu_penghasilan": "Tidak Berpenghasilan",
        "ibu_nik": "3671010101820001",
        "wali_nama": None, "wali_pendidikan": "Tidak sekolah", "wali_penghasilan": "< Rp1.000.000",
        "rombel": "7A", "no_registrasi_akta": "0001/KLU/JP/2013", "no_kk": "3671010101130001",
        "anak_ke": 1, "jml_saudara": 2, "berat_badan": 40, "tinggi_badan": 145,
        "lingkar_kepala": 52, "jarak_rumah": 1.5, "lintang": -6.1, "bujur": 106.6,
        "kebutuhan_khusus": "Tidak ada", "sekolah_asal": "SD NEGERI CONTOH 1",
    },
    {
        "nama": "CONTOH SISWA DUA", "nipd": "26260002", "jk": "P", "nisn": "3900000002",
        "tempat_lahir": "JAKARTA", "tanggal_lahir": "2013-03-08",
        "nik": "3173014803130002", "agama": "Islam", "alamat": "Jl. Contoh No. 2",
        "rt": "3", "rw": "4", "kelurahan": "Contoh Kelurahan", "kecamatan": "Kec. Contoh",
        "kode_pos": "15112", "jenis_tinggal": "Bersama orang tua", "transportasi": "Sepeda",
        "hp": "081200000002", "penerima_kps": "Ya", "no_kps": "0001234567890",
        "penerima_kip": "Ya", "nomor_kip": "KIP0000000002", "nama_kip": "CONTOH SISWA DUA",
        "layak_pip": "Ya", "alasan_layak_pip": "Penerima KIP",
        "ayah_nama": "AYAH CONTOH DUA", "ayah_tahun_lahir": 1978, "ayah_pendidikan": "SMA / sederajat",
        "ayah_pekerjaan": "Buruh Harian Lepas", "ayah_penghasilan": "< Rp1.000.000",
        "ibu_nama": "IBU CONTOH DUA", "ibu_tahun_lahir": 1981, "ibu_pendidikan": "SMA / sederajat",
        "ibu_pekerjaan": "Ibu Rumah Tangga", "ibu_penghasilan": "Tidak Berpenghasilan",
        "rombel": "7A", "bank": "BRI", "no_rekening": "000000000001",
        "rekening_atas_nama": "CONTOH SISWA DUA", "no_kk": "3173010101130002",
        "anak_ke": 2, "jml_saudara": 1, "berat_badan": 38, "tinggi_badan": 142,
        "kebutuhan_khusus": "Tidak ada", "sekolah_asal": "SD NEGERI CONTOH 2",
    },
]


def main() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    kolom = [spec.label for spec in STUDENT_FIELDS]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Daftar Peserta Didik"

    sheet["A1"] = "Daftar Peserta Didik (CONTOH TEMPLATE)"
    sheet["A1"].font = Font(size=14, bold=True)
    sheet["A2"] = "SEKOLAH CONTOH NEGERI 1"
    sheet["A3"] = "Kecamatan Kec. Contoh, Kabupaten Kota Contoh, Provinsi Prov. Contoh"
    sheet["A4"] = "Tanggal Unduh: 2026-09-16 08:00:00"

    judul_kolom = ["No", *kolom]
    for index, label in enumerate(judul_kolom, start=1):
        cell = sheet.cell(row=5, column=index, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, data in enumerate(CONTOH, start=6):
        sheet.cell(row=row_index, column=1, value=row_index - 5)
        for col_index, spec in enumerate(STUDENT_FIELDS, start=2):
            sheet.cell(row=row_index, column=col_index, value=data.get(spec.key))

    sheet.freeze_panes = "C6"
    sheet.column_dimensions["A"].width = 5
    sheet.column_dimensions["B"].width = 28
    for col_index in range(3, len(judul_kolom) + 1):
        sheet.column_dimensions[get_column_letter(col_index)].width = 18

    TUJUAN.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(TUJUAN)
    print(f"Template dibuat: {TUJUAN}")
    print(f"  {len(judul_kolom)} kolom, {len(CONTOH)} baris contoh data fiktif")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
