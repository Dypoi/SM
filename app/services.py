"""Lapisan layanan (business logic) SIMSEK.

Semua akses data terpusat di sini supaya router HTTP tetap tipis dan logika
mudah diuji maupun dipanggil ulang oleh bot Dapodik di tahap berikutnya.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from . import config, db
from .dapodik import (
    FIELD_BY_KEY,
    STUDENT_FIELDS,
    ParsedSpreadsheet,
    parse_sheet,
    summarize_issues,
)
from .readers import SheetData, format_label, read_spreadsheet

log = logging.getLogger("simsek.services")

# --------------------------------------------------------------------------- #
# Kolom yang boleh ditulis ke tabel students
# --------------------------------------------------------------------------- #
STUDENT_COLUMNS: tuple[str, ...] = tuple(spec.key for spec in STUDENT_FIELDS) + (
    "is_kps", "is_kip", "is_layak_pip", "status", "catatan",
)
STUDENT_WRITABLE = set(STUDENT_COLUMNS)

STATUS_OPTIONS = ("Aktif", "Lulus", "Mutasi", "Keluar", "Non-aktif")
AGAMA_OPTIONS = ("Islam", "Kristen", "Katholik", "Hindu", "Budha", "Konghucu", "Kepercayaan")

# Field yang dianggap penting untuk kelengkapan data (dipakai halaman Kualitas Data).
FIELD_WAJIB = ("nama", "nisn", "jk", "tempat_lahir", "tanggal_lahir", "nik",
               "alamat", "kelurahan", "kecamatan", "agama", "rombel")
FIELD_PENTING = ("nipd", "no_kk", "hp", "ayah_nama", "ayah_pekerjaan", "ibu_nama",
                 "ibu_pekerjaan", "no_registrasi_akta")


# =========================================================================== #
# PENGATURAN
# =========================================================================== #
def get_settings() -> dict[str, str]:
    return {row["key"]: row["value"] for row in db.query_all("SELECT key, value FROM settings")}


def get_setting(key: str, default: str | None = None) -> str | None:
    row = db.query_one("SELECT value FROM settings WHERE key = ?", (key,))
    if row is None or row["value"] is None:
        return default
    return row["value"]


def set_setting(key: str, value: Any) -> None:
    db.execute(
        """
        INSERT INTO settings(key, value, updated_at) VALUES(?, ?, datetime('now','localtime'))
        ON CONFLICT(key) DO UPDATE SET value = excluded.value, updated_at = excluded.updated_at
        """,
        (key, None if value is None else str(value)),
    )


def update_settings(values: dict[str, Any]) -> None:
    with db.transaction() as conn:
        for key, value in values.items():
            conn.execute(
                """
                INSERT INTO settings(key, value, updated_at) VALUES(?, ?, datetime('now','localtime'))
                ON CONFLICT(key) DO UPDATE SET value = excluded.value, updated_at = excluded.updated_at
                """,
                (key, None if value is None else str(value)),
            )


def school_profile() -> dict[str, str]:
    data = get_settings()
    return {
        "nama": data.get("school_nama") or "Sekolah",
        "npsn": data.get("school_npsn") or "",
        "alamat": data.get("school_alamat") or "",
        "kecamatan": data.get("school_kecamatan") or "",
        "kabupaten": data.get("school_kabupaten") or "",
        "provinsi": data.get("school_provinsi") or "",
        "kepala": data.get("school_kepala") or "",
        "email": data.get("school_email") or "",
        "telepon": data.get("school_telepon") or "",
        "tahun_ajaran": data.get("tahun_ajaran") or config.TAHUN_AJARAN_DEFAULT,
        "semester": data.get("semester") or config.SEMESTER_DEFAULT,
    }


# =========================================================================== #
# AUDIT
# =========================================================================== #
def log_audit(actor: str | None, role: str | None, aksi: str,
              entitas: str | None = None, entitas_id: Any = None,
              detail: str | None = None, ip: str | None = None) -> None:
    db.execute(
        """
        INSERT INTO audit_log(actor, role, aksi, entitas, entitas_id, detail, ip)
        VALUES(?,?,?,?,?,?,?)
        """,
        (actor, role, aksi, entitas, None if entitas_id is None else str(entitas_id), detail, ip),
    )


def recent_audit(limit: int = 20) -> list[dict[str, Any]]:
    return db.rows_to_dicts(
        db.query_all("SELECT * FROM audit_log ORDER BY id DESC LIMIT ?", (limit,))
    )


# =========================================================================== #
# IMPOR BERKAS
# =========================================================================== #
@dataclass
class StoredFile:
    path: Path
    filename: str
    size: int
    extension: str


def save_upload(filename: str, content: bytes) -> StoredFile:
    """Simpan berkas unggahan ke folder data (nama unik, aman dari path traversal)."""
    config.ensure_dirs()
    safe_name = re.sub(r"[^\w.\- ]+", "_", Path(filename).name).strip() or "unggahan"
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    target = config.UPLOAD_DIR / f"{stamp}__{safe_name}"
    counter = 1
    while target.exists():
        target = config.UPLOAD_DIR / f"{stamp}-{counter}__{safe_name}"
        counter += 1
    target.write_bytes(content)
    return StoredFile(
        path=target,
        filename=filename,
        size=len(content),
        extension=target.suffix.lower(),
    )


def create_import(
    stored: StoredFile,
    *,
    sheet: SheetData,
    parsed: ParsedSpreadsheet,
    fmt: str,
    actor: str | None,
    mode: str = "upsert",
) -> int:
    summary = summarize_issues(parsed.issues)
    status = "draft"
    import_id = db.insert_returning_id(
        """
        INSERT INTO imports(
            filename, stored_path, file_ext, file_size, file_format, sheet_name,
            header_row, mapping_json, meta_json, mode, status, rows_total,
            detected_kind, message, created_by
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            stored.filename,
            str(stored.path),
            stored.extension,
            stored.size,
            fmt,
            sheet.name,
            parsed.header_row,
            json.dumps(parsed.mapping, ensure_ascii=False),
            json.dumps({"metadata": parsed.metadata, **summary}, ensure_ascii=False),
            mode,
            status,
            len(parsed.records),
            parsed.kind,
            f"{len(parsed.records)} baris data terdeteksi, {summary['error']} error, {summary['warning']} peringatan.",
            actor,
        ),
    )
    for issue in parsed.issues[:2000]:
        db.execute(
            "INSERT INTO import_issues(import_id, row_number, level, field, message) VALUES(?,?,?,?,?)",
            (import_id, issue.row_number, issue.level, issue.field, issue.message),
        )
    return import_id


def list_imports(limit: int = 50) -> list[dict[str, Any]]:
    return db.rows_to_dicts(
        db.query_all("SELECT * FROM imports ORDER BY id DESC LIMIT ?", (limit,))
    )


def get_import(import_id: int) -> dict[str, Any] | None:
    row = db.query_one("SELECT * FROM imports WHERE id = ?", (import_id,))
    return db.row_to_dict(row)


def import_issues(import_id: int, level: str | None = None, limit: int = 500) -> list[dict[str, Any]]:
    if level:
        rows = db.query_all(
            "SELECT * FROM import_issues WHERE import_id = ? AND level = ? ORDER BY row_number LIMIT ?",
            (import_id, level, limit),
        )
    else:
        rows = db.query_all(
            "SELECT * FROM import_issues WHERE import_id = ? ORDER BY row_number LIMIT ?",
            (import_id, limit),
        )
    return db.rows_to_dicts(rows)


def count_import_issues(import_id: int) -> dict[str, int]:
    rows = db.query_all(
        "SELECT level, COUNT(*) AS jumlah FROM import_issues WHERE import_id = ? GROUP BY level",
        (import_id,),
    )
    result = {"error": 0, "warning": 0, "info": 0}
    for row in rows:
        result[row["level"]] = row["jumlah"]
    return result


def delete_import(import_id: int) -> None:
    record = get_import(import_id)
    if record and record.get("stored_path"):
        try:
            Path(record["stored_path"]).unlink(missing_ok=True)
        except OSError:  # pragma: no cover
            pass
    db.execute("DELETE FROM imports WHERE id = ?", (import_id,))


@dataclass
class ImportResult:
    imported: int = 0
    updated: int = 0
    skipped: int = 0
    failed: int = 0
    total: int = 0
    issues: list[str] = None  # type: ignore[assignment]

    def __post_init__(self) -> None:
        if self.issues is None:
            self.issues = []


def import_students(
    parsed: ParsedSpreadsheet,
    *,
    import_id: int | None = None,
    mode: str = "upsert",
    actor: str | None = None,
    source_file: str | None = None,
    only_valid: bool = True,
    column_map: dict[int, str] | None = None,
) -> ImportResult:
    """Simpan hasil parsing ke tabel ``students``.

    mode:
        ``insert``  - hanya tambah siswa baru, NISN yang sudah ada dilewati
        ``upsert``  - tambah baru / perbarui data yang sudah ada (default)
        ``replace`` - perbarui seluruh field yang tersedia, termasuk menimpa manual
        ``dry``     - hanya validasi, tidak menulis apa pun
    """
    result = ImportResult(total=len(parsed.records))
    if mode == "dry":
        return result

    with db.transaction() as conn:
        for record in parsed.records:
            if only_valid and record.has_error:
                result.failed += 1
                result.issues.append(
                    f"Baris {record.row_number}: dilewati karena {record.issues[0].message}"
                    if record.issues else f"Baris {record.row_number}: dilewati (data tidak valid)"
                )
                continue

            values = {key: value for key, value in record.values.items() if key in STUDENT_WRITABLE}
            if not values.get("nama"):
                result.failed += 1
                continue

            nisn = values.get("nisn")
            existing = None
            if nisn:
                existing = conn.execute("SELECT id FROM students WHERE nisn = ?", (nisn,)).fetchone()
            if existing is None and values.get("nipd"):
                existing = conn.execute(
                    "SELECT id FROM students WHERE nipd = ? AND (nisn IS NULL OR nisn = '')",
                    (values["nipd"],),
                ).fetchone()

            if existing is None:
                if mode == "insert" or mode in {"upsert", "replace"}:
                    columns = list(values.keys()) + ["source_file", "import_id"]
                    placeholders = ", ".join("?" for _ in columns)
                    conn.execute(
                        f"INSERT INTO students({', '.join(columns)}) VALUES({placeholders})",
                        [*values.values(), source_file, import_id],
                    )
                    result.imported += 1
                else:
                    result.skipped += 1
                continue

            student_id = existing["id"]
            if mode == "insert":
                result.skipped += 1
                continue

            current = dict(conn.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone())
            changes = []
            updates: dict[str, Any] = {}
            for key, new_value in values.items():
                old_value = current.get(key)
                if _same_value(old_value, new_value):
                    continue
                if mode == "upsert" and old_value not in (None, "") and key in {"nik", "no_kk", "nama"}:
                    # pada mode upsert, identitas kunci yang sudah terisi tidak ditimpa
                    continue
                updates[key] = new_value
                changes.append((key, old_value, new_value))

            if not updates:
                result.skipped += 1
                continue

            updates["source_file"] = source_file
            updates["import_id"] = import_id
            set_clause = ", ".join(f"{key} = ?" for key in updates)
            conn.execute(
                f"UPDATE students SET {set_clause}, updated_at = datetime('now','localtime') WHERE id = ?",
                [*updates.values(), student_id],
            )
            for field, old_value, new_value in changes:
                conn.execute(
                    """
                    INSERT INTO data_changes(student_id, nisn, field, old_value, new_value, source, actor)
                    VALUES(?,?,?,?,?,?,?)
                    """,
                    (student_id, nisn, field, _text(old_value), _text(new_value), "impor", actor),
                )
            result.updated += 1

    if import_id is not None:
        status = "sukses"
        if result.failed and result.imported + result.updated == 0:
            status = "gagal"
        elif result.failed:
            status = "sebagian"
        db.execute(
            """
            UPDATE imports SET status = ?, rows_imported = ?, rows_updated = ?,
                   rows_skipped = ?, rows_failed = ?, finished_at = datetime('now','localtime')
             WHERE id = ?
            """,
            (status, result.imported, result.updated, result.skipped, result.failed, import_id),
        )
    return result


def _text(value: Any) -> str | None:
    return None if value is None else str(value)


def _same_value(a: Any, b: Any) -> bool:
    if a is None and b is None:
        return True
    if isinstance(a, str) and isinstance(b, str):
        return a.strip().casefold() == b.strip().casefold()
    if a is None or b is None:
        return False
    try:
        return float(a) == float(b)
    except (TypeError, ValueError):
        return str(a).strip() == str(b).strip()


def read_and_parse(path: str | Path, ext: str | None = None, sheet_index: int = 0,
                   column_map: dict[int, str] | None = None) -> tuple[SheetData, ParsedSpreadsheet, str]:
    """Helper: baca berkas lalu urai isinya (dipakai halaman impor)."""
    sheets = read_spreadsheet(path, ext)
    sheet = sheets[min(sheet_index, len(sheets) - 1)]
    parsed = parse_sheet(sheet, column_map=column_map) if column_map else parse_sheet(sheet)
    extension = Path(path).suffix.lower()
    return sheet, parsed, format_label(extension)


# =========================================================================== #
# PESERTA DIDIK
# =========================================================================== #
SORTABLE_COLUMNS = {
    "nama": "nama",
    "nisn": "nisn",
    "nipd": "nipd",
    "rombel": "rombel",
    "tingkat": "tingkat",
    "tanggal_lahir": "tanggal_lahir",
    "kelurahan": "kelurahan",
    "updated_at": "updated_at",
}


@dataclass
class StudentFilter:
    q: str = ""
    rombel: str = ""
    tingkat: str = ""
    jk: str = ""
    agama: str = ""
    status: str = ""
    kelurahan: str = ""
    sort: str = "nama"
    direction: str = "asc"
    incomplete_only: bool = False

    def where_clause(self) -> tuple[str, list[Any]]:
        clauses: list[str] = []
        params: list[Any] = []

        if self.q:
            like = f"%{self.q.strip()}%"
            clauses.append(
                "(s.nama LIKE ? OR s.nisn LIKE ? OR s.nipd LIKE ? OR s.nik LIKE ? "
                "OR s.no_kk LIKE ? OR s.rombel LIKE ? OR s.ayah_nama LIKE ? OR s.ibu_nama LIKE ?)"
            )
            params.extend([like] * 8)
        if self.rombel:
            clauses.append("s.rombel = ?")
            params.append(self.rombel)
        if self.tingkat:
            clauses.append("s.tingkat = ?")
            params.append(self.tingkat)
        if self.jk:
            clauses.append("s.jk = ?")
            params.append(self.jk)
        if self.agama:
            clauses.append("s.agama = ?")
            params.append(self.agama)
        if self.status:
            clauses.append("s.status = ?")
            params.append(self.status)
        if self.kelurahan:
            clauses.append("s.kelurahan = ?")
            params.append(self.kelurahan)
        if self.incomplete_only:
            checks = " OR ".join(f"s.{field} IS NULL OR TRIM(s.{field}) = ''" for field in FIELD_WAJIB)
            clauses.append(f"({checks})")

        return (" AND ".join(clauses) if clauses else "1=1"), params

    def order_clause(self) -> str:
        column = SORTABLE_COLUMNS.get(self.sort, "nama")
        direction = "DESC" if str(self.direction).lower() == "desc" else "ASC"
        return f"{column} COLLATE NOCASE {direction}, s.id {direction}"


def _pick_columns(include_detail: bool = False) -> str:
    base = (
        "s.id, s.nama, s.nisn, s.nipd, s.jk, s.rombel, s.tingkat, s.tempat_lahir, "
        "s.tanggal_lahir, s.nik, s.agama, s.kelurahan, s.kecamatan, s.hp, s.status, "
        "s.is_kip, s.is_kps, s.is_layak_pip, s.updated_at, s.alamat"
    )
    return base + (", s.*" if include_detail else "")


def list_students(
    filters: StudentFilter,
    page: int = 1,
    per_page: int | None = None,
) -> tuple[list[dict[str, Any]], int]:
    per_page = per_page or config.ROWS_PER_PAGE
    where, params = filters.where_clause()
    total = int(db.query_value(f"SELECT COUNT(*) FROM students s WHERE {where}", params, 0) or 0)
    offset = max(0, (page - 1) * per_page)
    rows = db.query_all(
        f"""
        SELECT {_pick_columns()} FROM students s
        WHERE {where}
        ORDER BY {filters.order_clause()}
        LIMIT ? OFFSET ?
        """,
        [*params, per_page, offset],
    )
    return db.rows_to_dicts(rows), total


def get_student(student_id: int) -> dict[str, Any] | None:
    return db.row_to_dict(db.query_one("SELECT * FROM students WHERE id = ?", (student_id,)))


def get_student_by_nisn(nisn: str) -> dict[str, Any] | None:
    return db.row_to_dict(db.query_one("SELECT * FROM students WHERE nisn = ?", (str(nisn).strip(),)))


def student_exists(nisn: str | None, nipd: str | None = None,
                   exclude_id: int | None = None) -> dict[str, Any] | None:
    """Cek duplikat NISN/NIPD sebelum simpan manual."""
    params: list[Any] = []
    clauses: list[str] = []
    if nisn:
        clauses.append("nisn = ?")
        params.append(nisn)
    if nipd:
        clauses.append("nipd = ?")
        params.append(nipd)
    if not clauses:
        return None
    sql = f"SELECT id, nama, nisn, nipd FROM students WHERE ({' OR '.join(clauses)})"
    if exclude_id:
        sql += " AND id <> ?"
        params.append(exclude_id)
    return db.row_to_dict(db.query_one(sql, params))


def create_student(data: dict[str, Any], actor: str | None = None) -> int:
    values = {key: value for key, value in data.items() if key in STUDENT_WRITABLE}
    if not values.get("nama"):
        raise ValueError("Nama siswa wajib diisi.")
    if not values.get("tingkat") and values.get("rombel"):
        from .dapodik import derive_tingkat

        values["tingkat"] = derive_tingkat(values["rombel"])
    columns = list(values.keys())
    placeholders = ", ".join("?" for _ in columns)
    student_id = db.insert_returning_id(
        f"INSERT INTO students({', '.join(columns)}) VALUES({placeholders})",
        [values[key] for key in columns],
    )
    for key, value in values.items():
        db.execute(
            "INSERT INTO data_changes(student_id, nisn, field, new_value, source, actor) VALUES(?,?,?,?,?,?)",
            (student_id, values.get("nisn"), key, _text(value), "manual", actor),
        )
    log_audit(actor, None, "tambah_siswa", "students", student_id, values.get("nama"))
    return student_id


def update_student(student_id: int, data: dict[str, Any],
                   actor: str | None = None, source: str = "manual") -> list[tuple[str, Any, Any]]:
    """Perbarui siswa dan catat setiap perubahan field (untuk audit & bot Dapodik)."""
    current = get_student(student_id)
    if current is None:
        raise ValueError("Siswa tidak ditemukan.")

    updates: dict[str, Any] = {}
    changes: list[tuple[str, Any, Any]] = []
    for key, new_value in data.items():
        if key not in STUDENT_WRITABLE or key in {"status"} and new_value == "":
            continue
        old_value = current.get(key)
        if _same_value(old_value, new_value):
            continue
        updates[key] = new_value
        changes.append((key, old_value, new_value))

    if not updates:
        return []
    if "rombel" in updates and "tingkat" not in updates:
        from .dapodik import derive_tingkat

        updates["tingkat"] = derive_tingkat(updates["rombel"])

    set_clause = ", ".join(f"{key} = ?" for key in updates)
    with db.transaction() as conn:
        conn.execute(
            f"UPDATE students SET {set_clause}, updated_at = datetime('now','localtime') WHERE id = ?",
            [*updates.values(), student_id],
        )
        for field, old_value, new_value in changes:
            conn.execute(
                """
                INSERT INTO data_changes(student_id, nisn, field, old_value, new_value, source, actor)
                VALUES(?,?,?,?,?,?,?)
                """,
                (student_id, current.get("nisn"), field, _text(old_value), _text(new_value), source, actor),
            )
    log_audit(actor, None, "ubah_siswa", "students", student_id, ", ".join(f"{f}" for f, _, _ in changes))
    return changes


def delete_student(student_id: int, actor: str | None = None) -> None:
    student = get_student(student_id)
    db.execute("DELETE FROM students WHERE id = ?", (student_id,))
    if student:
        log_audit(actor, None, "hapus_siswa", "students", student_id, student.get("nama"))


def student_changes(student_id: int, limit: int = 50) -> list[dict[str, Any]]:
    return db.rows_to_dicts(
        db.query_all(
            "SELECT * FROM data_changes WHERE student_id = ? ORDER BY id DESC LIMIT ?",
            (student_id, limit),
        )
    )


def distinct_values(column: str, limit: int = 500) -> list[str]:
    if column not in {"rombel", "tingkat", "jk", "agama", "kelurahan", "kecamatan", "status"}:
        raise ValueError(f"Kolom '{column}' tidak boleh dipakai untuk filter.")
    rows = db.query_all(
        f"SELECT DISTINCT {column} AS v FROM students WHERE {column} IS NOT NULL AND TRIM({column}) <> '' "
        f"ORDER BY {column} COLLATE NOCASE LIMIT ?",
        (limit,),
    )
    return [row["v"] for row in rows]


def student_stats() -> dict[str, Any]:
    total = int(db.query_value("SELECT COUNT(*) FROM students") or 0)
    laki = int(db.query_value("SELECT COUNT(*) FROM students WHERE jk = 'L'") or 0)
    perempuan = int(db.query_value("SELECT COUNT(*) FROM students WHERE jk = 'P'") or 0)
    rombel_count = int(db.query_value("SELECT COUNT(DISTINCT rombel) FROM students WHERE rombel IS NOT NULL AND rombel <> ''") or 0)
    kip = int(db.query_value("SELECT COUNT(*) FROM students WHERE is_kip = 1") or 0)
    kps = int(db.query_value("SELECT COUNT(*) FROM students WHERE is_kps = 1") or 0)
    pip = int(db.query_value("SELECT COUNT(*) FROM students WHERE is_layak_pip = 1") or 0)
    khusus = int(db.query_value("SELECT COUNT(*) FROM students WHERE kebutuhan_khusus IS NOT NULL AND kebutuhan_khusus <> 'Tidak ada'") or 0)
    return {
        "total": total,
        "laki_laki": laki,
        "perempuan": perempuan,
        "rombel": rombel_count,
        "penerima_kip": kip,
        "penerima_kps": kps,
        "layak_pip": pip,
        "kebutuhan_khusus": khusus,
    }


def stats_by_rombel() -> list[dict[str, Any]]:
    rows = db.query_all(
        """
        SELECT rombel,
               COUNT(*) AS jumlah,
               SUM(CASE WHEN jk = 'L' THEN 1 ELSE 0 END) AS laki_laki,
               SUM(CASE WHEN jk = 'P' THEN 1 ELSE 0 END) AS perempuan,
               MAX(tingkat) AS tingkat
          FROM students
         WHERE rombel IS NOT NULL AND TRIM(rombel) <> ''
         GROUP BY rombel
         ORDER BY MIN(tingkat), rombel COLLATE NOCASE
        """
    )
    return db.rows_to_dicts(rows)


def stats_by_tingkat() -> list[dict[str, Any]]:
    rows = db.query_all(
        """
        SELECT COALESCE(NULLIF(tingkat,''),'-') AS tingkat,
               COUNT(*) AS jumlah,
               COUNT(DISTINCT rombel) AS rombel,
               SUM(CASE WHEN jk = 'L' THEN 1 ELSE 0 END) AS laki_laki,
               SUM(CASE WHEN jk = 'P' THEN 1 ELSE 0 END) AS perempuan
          FROM students
         GROUP BY COALESCE(NULLIF(tingkat,''),'-')
         ORDER BY tingkat
        """
    )
    return db.rows_to_dicts(rows)


def stats_by(column: str, limit: int = 12) -> list[dict[str, Any]]:
    if column not in {"agama", "kelurahan", "kecamatan", "transportasi", "kebutuhan_khusus", "jenis_tinggal", "sekolah_asal"}:
        raise ValueError("Kolom statistik tidak diizinkan.")
    rows = db.query_all(
        f"""
        SELECT COALESCE(NULLIF({column},''),'(belum diisi)') AS label, COUNT(*) AS jumlah
          FROM students
         GROUP BY COALESCE(NULLIF({column},''),'(belum diisi)')
         ORDER BY jumlah DESC
         LIMIT ?
        """,
        (limit,),
    )
    return db.rows_to_dicts(rows)


def data_quality() -> dict[str, Any]:
    """Ringkas kelengkapan data siswa per field — bahan perbaikan sebelum sinkron Dapodik."""
    total = int(db.query_value("SELECT COUNT(*) FROM students") or 0)
    items: list[dict[str, Any]] = []
    for spec in STUDENT_FIELDS:
        filled = int(
            db.query_value(
                f"SELECT COUNT(*) FROM students WHERE {spec.key} IS NOT NULL AND TRIM(CAST({spec.key} AS TEXT)) <> ''"
            )
            or 0
        )
        items.append(
            {
                "key": spec.key,
                "label": spec.label,
                "group": spec.group,
                "wajib": spec.key in FIELD_WAJIB,
                "terisi": filled,
                "kosong": total - filled,
                "persen": round(filled / total * 100, 1) if total else 0.0,
            }
        )

    # Siswa dengan data wajib belum lengkap
    checks = " OR ".join(f"s.{field} IS NULL OR TRIM(CAST(s.{field} AS TEXT)) = ''" for field in FIELD_WAJIB)
    belum_lengkap = int(db.query_value(f"SELECT COUNT(*) FROM students s WHERE {checks}") or 0)

    # Temuan spesifik ala Dapodik
    temuan = [
        {
            "kode": "NISN_TIDAK_VALID",
            "label": "NISN bukan 10 digit angka",
            "jumlah": int(db.query_value(
                "SELECT COUNT(*) FROM students WHERE nisn IS NULL OR LENGTH(TRIM(nisn)) <> 10 OR TRIM(nisn) GLOB '*[^0-9]*'"
            ) or 0),
            "field": "nisn",
        },
        {
            "kode": "NIK_TIDAK_VALID",
            "label": "NIK bukan 16 digit angka",
            "jumlah": int(db.query_value(
                "SELECT COUNT(*) FROM students WHERE nik IS NULL OR TRIM(nik) = '' OR LENGTH(TRIM(nik)) <> 16 OR TRIM(nik) GLOB '*[^0-9]*'"
            ) or 0),
            "field": "nik",
        },
        {
            "kode": "TANGGAL_LAHIR_KOSONG",
            "label": "Tanggal lahir belum diisi",
            "jumlah": int(db.query_value("SELECT COUNT(*) FROM students WHERE tanggal_lahir IS NULL OR TRIM(tanggal_lahir) = ''") or 0),
            "field": "tanggal_lahir",
        },
        {
            "kode": "ROMBEL_KOSONG",
            "label": "Rombel belum diisi",
            "jumlah": int(db.query_value("SELECT COUNT(*) FROM students WHERE rombel IS NULL OR TRIM(rombel) = ''") or 0),
            "field": "rombel",
        },
        {
            "kode": "NO_KK_KOSONG",
            "label": "Nomor Kartu Keluarga belum diisi",
            "jumlah": int(db.query_value("SELECT COUNT(*) FROM students WHERE no_kk IS NULL OR TRIM(no_kk) = ''") or 0),
            "field": "no_kk",
        },
        {
            "kode": "NISN_GANDA",
            "label": "NISN ganda",
            "jumlah": int(db.query_value(
                "SELECT COUNT(*) FROM (SELECT nisn FROM students WHERE nisn IS NOT NULL AND TRIM(nisn) <> '' GROUP BY nisn HAVING COUNT(*) > 1)"
            ) or 0),
            "field": "nisn",
        },
        {
            "kode": "ALAMAT_KOSONG",
            "label": "Alamat belum diisi",
            "jumlah": int(db.query_value("SELECT COUNT(*) FROM students WHERE alamat IS NULL OR TRIM(alamat) = ''") or 0),
            "field": "alamat",
        },
        {
            "kode": "IBU_KOSONG",
            "label": "Data ibu belum diisi",
            "jumlah": int(db.query_value("SELECT COUNT(*) FROM students WHERE ibu_nama IS NULL OR TRIM(ibu_nama) = ''") or 0),
            "field": "ibu_nama",
        },
    ]
    return {
        "total": total,
        "fields": items,
        "belum_lengkap": belum_lengkap,
        "siap_sinkron": max(0, total - belum_lengkap),
        "temuan": temuan,
    }


def students_with_duplicate_nisn() -> list[dict[str, Any]]:
    return db.rows_to_dicts(
        db.query_all(
            """
            SELECT nisn, COUNT(*) AS jumlah, GROUP_CONCAT(nama, ' | ') AS nama
              FROM students
             WHERE nisn IS NOT NULL AND TRIM(nisn) <> ''
             GROUP BY nisn HAVING COUNT(*) > 1
             ORDER BY jumlah DESC
            """
        )
    )


# --------------------------------------------------------------------------- #
# Ekspor
# --------------------------------------------------------------------------- #
def _export_field_keys() -> list[str]:
    return [spec.key for spec in STUDENT_FIELDS]


def export_students_csv(filters: StudentFilter) -> bytes:
    where, params = filters.where_clause()
    rows = db.query_all(f"SELECT * FROM students s WHERE {where} ORDER BY {filters.order_clause()}", params)
    columns = _export_field_keys()
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow(["No", *[FIELD_BY_KEY[key].label for key in columns], "Status"])
    for index, row in enumerate(rows, start=1):
        writer.writerow([index, *[row[key] if key in row.keys() else "" for key in columns], row["status"]])
    return buffer.getvalue().encode("utf-8-sig")


def export_students_xlsx(filters: StudentFilter) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    where, params = filters.where_clause()
    rows = db.query_all(f"SELECT * FROM students s WHERE {where} ORDER BY {filters.order_clause()}", params)
    columns = _export_field_keys()

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Daftar Peserta Didik"
    profile = school_profile()

    sheet["A1"] = "Daftar Peserta Didik"
    sheet["A1"].font = Font(size=14, bold=True)
    sheet["A2"] = profile["nama"]
    sheet["A3"] = f"Kecamatan {profile['kecamatan']}, {profile['kabupaten']}, {profile['provinsi']}".strip(", ")
    sheet["A4"] = f"Tahun Ajaran {profile['tahun_ajaran']} | Diekspor dari SIMSEK pada {dt.datetime.now():%Y-%m-%d %H:%M}"

    header = ["No", *[FIELD_BY_KEY[key].label for key in columns], "Status"]
    for col_index, label in enumerate(header, start=1):
        cell = sheet.cell(row=5, column=col_index, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, row in enumerate(rows, start=6):
        sheet.cell(row=row_index, column=1, value=row_index - 5)
        for col_index, key in enumerate(columns, start=2):
            value = row[key] if key in row.keys() else None
            sheet.cell(row=row_index, column=col_index, value=value)

    sheet.freeze_panes = "C6"
    sheet.auto_filter.ref = f"A5:{get_column_letter(len(header))}{max(5, 5 + len(rows))}"
    sheet.column_dimensions["A"].width = 5
    sheet.column_dimensions["B"].width = 30

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


# =========================================================================== #
# EKSTRAKURIKULER
# =========================================================================== #
EKSKUL_KATEGORI = ("Olahraga", "Seni", "Akademik", "Keagamaan", "Kepemimpinan", "Teknologi", "Lainnya")
HARI_OPTIONS = ("Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu")


def list_ekskul(aktif_only: bool = False) -> list[dict[str, Any]]:
    sql = """
        SELECT e.*, (SELECT COUNT(*) FROM ekskul_members m WHERE m.ekskul_id = e.id AND m.status='aktif') AS jumlah_anggota
          FROM extracurriculars e
    """
    if aktif_only:
        sql += " WHERE e.aktif = 1"
    sql += " ORDER BY e.aktif DESC, e.nama COLLATE NOCASE"
    return db.rows_to_dicts(db.query_all(sql))


def get_ekskul(ekskul_id: int) -> dict[str, Any] | None:
    return db.row_to_dict(db.query_one("SELECT * FROM extracurriculars WHERE id = ?", (ekskul_id,)))


def save_ekskul(data: dict[str, Any], ekskul_id: int | None = None, actor: str | None = None) -> int:
    fields = ("kode", "nama", "kategori", "pembina", "hari", "jam_mulai", "jam_selesai",
              "tempat", "kuota", "deskripsi", "aktif")
    values = {key: data.get(key) for key in fields if key in data}
    if not values.get("nama"):
        raise ValueError("Nama ekstrakurikuler wajib diisi.")

    if ekskul_id:
        set_clause = ", ".join(f"{key} = ?" for key in values)
        db.execute(
            f"UPDATE extracurriculars SET {set_clause}, updated_at = datetime('now','localtime') WHERE id = ?",
            [*values.values(), ekskul_id],
        )
        log_audit(actor, None, "ubah_ekskul", "extracurriculars", ekskul_id, values.get("nama"))
        return ekskul_id

    columns = list(values.keys())
    placeholders = ", ".join("?" for _ in columns)
    new_id = db.insert_returning_id(
        f"INSERT INTO extracurriculars({', '.join(columns)}) VALUES({placeholders})",
        [values[key] for key in columns],
    )
    log_audit(actor, None, "tambah_ekskul", "extracurriculars", new_id, values.get("nama"))
    return new_id


def delete_ekskul(ekskul_id: int, actor: str | None = None) -> None:
    ekskul = get_ekskul(ekskul_id)
    db.execute("DELETE FROM extracurriculars WHERE id = ?", (ekskul_id,))
    if ekskul:
        log_audit(actor, None, "hapus_ekskul", "extracurriculars", ekskul_id, ekskul.get("nama"))


def ekskul_members(ekskul_id: int) -> list[dict[str, Any]]:
    return db.rows_to_dicts(
        db.query_all(
            """
            SELECT m.*, s.nama, s.nisn, s.rombel, s.tingkat, s.jk
              FROM ekskul_members m
              JOIN students s ON s.id = m.student_id
             WHERE m.ekskul_id = ?
             ORDER BY m.status, s.rombel, s.nama COLLATE NOCASE
            """,
            (ekskul_id,),
        )
    )


def add_ekskul_member(ekskul_id: int, student_id: int, *, jabatan: str = "Anggota",
                      tahun_ajaran: str | None = None, semester: str | None = None,
                      actor: str | None = None) -> int | None:
    profile = school_profile()
    tahun_ajaran = tahun_ajaran or profile["tahun_ajaran"]
    semester = semester or profile["semester"]
    existing = db.query_one(
        "SELECT id FROM ekskul_members WHERE ekskul_id = ? AND student_id = ? AND tahun_ajaran = ?",
        (ekskul_id, student_id, tahun_ajaran),
    )
    if existing:
        return None
    member_id = db.insert_returning_id(
        """
        INSERT INTO ekskul_members(ekskul_id, student_id, jabatan, tahun_ajaran, semester)
        VALUES(?,?,?,?,?)
        """,
        (ekskul_id, student_id, jabatan, tahun_ajaran, semester),
    )
    log_audit(actor, None, "tambah_anggota_ekskul", "ekskul_members", member_id, f"ekskul={ekskul_id}")
    return member_id


def remove_ekskul_member(member_id: int, actor: str | None = None) -> None:
    db.execute("DELETE FROM ekskul_members WHERE id = ?", (member_id,))
    log_audit(actor, None, "hapus_anggota_ekskul", "ekskul_members", member_id)


def update_ekskul_member(member_id: int, data: dict[str, Any], actor: str | None = None) -> None:
    allowed = {key: value for key, value in data.items() if key in {"jabatan", "nilai", "predikat", "status"}}
    if not allowed:
        return
    set_clause = ", ".join(f"{key} = ?" for key in allowed)
    db.execute(f"UPDATE ekskul_members SET {set_clause} WHERE id = ?", [*allowed.values(), member_id])
    log_audit(actor, None, "ubah_anggota_ekskul", "ekskul_members", member_id)


def student_ekskul(student_id: int) -> list[dict[str, Any]]:
    return db.rows_to_dicts(
        db.query_all(
            """
            SELECT e.nama, e.kategori, e.hari, e.jam_mulai, e.jam_selesai, e.tempat, e.pembina,
                   m.jabatan, m.status, m.predikat, m.nilai, m.tahun_ajaran
              FROM ekskul_members m
              JOIN extracurriculars e ON e.id = m.ekskul_id
             WHERE m.student_id = ?
             ORDER BY e.nama COLLATE NOCASE
            """,
            (student_id,),
        )
    )


def ekskul_stats() -> dict[str, Any]:
    total = int(db.query_value("SELECT COUNT(*) FROM extracurriculars") or 0)
    aktif = int(db.query_value("SELECT COUNT(*) FROM extracurriculars WHERE aktif = 1") or 0)
    anggota = int(db.query_value("SELECT COUNT(*) FROM ekskul_members WHERE status = 'aktif'") or 0)
    siswa_ikut = int(db.query_value("SELECT COUNT(DISTINCT student_id) FROM ekskul_members") or 0)
    return {"total": total, "aktif": aktif, "anggota": anggota, "siswa_ikut": siswa_ikut}


def ekskul_by_kategori() -> list[dict[str, Any]]:
    return db.rows_to_dicts(
        db.query_all(
            """
            SELECT COALESCE(NULLIF(kategori,''),'Lainnya') AS label, COUNT(*) AS jumlah
              FROM extracurriculars GROUP BY COALESCE(NULLIF(kategori,''),'Lainnya')
             ORDER BY jumlah DESC
            """
        )
    )


# =========================================================================== #
# BOT DAPODIK (kerangka kerja)
# =========================================================================== #
def list_dapodik_jobs(limit: int = 20) -> list[dict[str, Any]]:
    return db.rows_to_dicts(
        db.query_all("SELECT * FROM dapodik_jobs ORDER BY id DESC LIMIT ?", (limit,))
    )


def create_dapodik_job(jenis: str, mode: str, payload: dict[str, Any] | None = None,
                       actor: str | None = None) -> int:
    job_id = db.insert_returning_id(
        "INSERT INTO dapodik_jobs(jenis, mode, payload_json, created_by) VALUES(?,?,?,?)",
        (jenis, mode, json.dumps(payload or {}, ensure_ascii=False), actor),
    )
    log_audit(actor, None, "buat_job_dapodik", "dapodik_jobs", job_id, f"{jenis}/{mode}")
    return job_id


def dapodik_readiness() -> dict[str, Any]:
    """Ringkasan kesiapan data untuk sinkronisasi ke Dapodik."""
    quality = data_quality()
    return {
        "total": quality["total"],
        "siap": quality["siap_sinkron"],
        "perlu_perbaikan": quality["belum_lengkap"],
        "temuan": quality["temuan"],
        "api_key": get_setting("api_key"),
        "sync_aktif": get_setting("dapodik_sync_aktif", "0") == "1",
    }


# =========================================================================== #
# SEED DATA CONTOH
# =========================================================================== #
EKSKUL_SEED = [
    ("PKS", "Pasukan Khusus Sekolah", "Kepemimpinan", "Budiman, S.Pd", "Sabtu", "08:00", "10:00", "Lapangan"),
    ("PMR", "Palang Merah Remaja", "Kepemimpinan", "Siti Rohmah, S.Pd", "Jumat", "14:00", "16:00", "Ruang UKS"),
    ("FUTSAL", "Futsal", "Olahraga", "Agus Setiawan, S.Pd", "Selasa", "15:00", "17:00", "Lapangan Futsal"),
    ("BASKET", "Bola Basket", "Olahraga", "Rina Marlina, S.Pd", "Rabu", "15:00", "17:00", "Lapangan Basket"),
    ("PRAMUKA", "Pramuka", "Kepemimpinan", "Hendra Gunawan, S.Pd", "Jumat", "14:00", "16:30", "Lapangan"),
    ("ROHIS", "Rohani Islam", "Keagamaan", "Ustadz Ahmad Fauzi", "Kamis", "14:00", "15:30", "Musala"),
    ("ENGLISH", "English Club", "Akademik", "Dewi Lestari, S.Pd", "Senin", "14:00", "15:30", "Ruang 9A"),
    ("TARI", "Seni Tari", "Seni", "Wulan Sari, S.Pd", "Sabtu", "09:00", "11:00", "Aula"),
    ("BAND", "Band & Musik", "Seni", "Rizky Pratama, S.Pd", "Rabu", "14:00", "16:00", "Ruang Musik"),
    ("CODING", "Klub Coding", "Teknologi", "Bayu Nugroho, S.Kom", "Selasa", "14:00", "16:00", "Lab Komputer"),
]


def seed_ekskul_if_empty() -> int:
    existing = int(db.query_value("SELECT COUNT(*) FROM extracurriculars") or 0)
    if existing:
        return 0
    count = 0
    for kode, nama, kategori, pembina, hari, mulai, selesai, tempat in EKSKUL_SEED:
        db.execute(
            """
            INSERT INTO extracurriculars(kode, nama, kategori, pembina, hari, jam_mulai, jam_selesai, tempat, kuota, deskripsi)
            VALUES(?,?,?,?,?,?,?,?,?,?)
            """,
            (kode, nama, kategori, pembina, hari, mulai, selesai, tempat, 40, f"Ekstrakurikuler {nama}"),
        )
        count += 1
    return count


def find_sample_files() -> list[Path]:
    """Cari berkas contoh Dapodik di repo (untuk pengisian data awal)."""
    if not config.SAMPLE_DIR.exists():
        return []
    files: list[Path] = []
    for pattern in ("*.xlsx", "*.xls", "*.csv"):
        files.extend(sorted(config.SAMPLE_DIR.glob(pattern)))
    return files


def auto_seed_sample(nisn_example: dict[str, str] | None = None) -> dict[str, Any] | None:
    """Impor otomatis berkas contoh bila database masih kosong."""
    if not config.AUTO_SEED:
        return None
    total = int(db.query_value("SELECT COUNT(*) FROM students") or 0)
    if total:
        return None
    samples = find_sample_files()
    if not samples:
        return None
    path = samples[0]
    try:
        sheets = read_spreadsheet(path, path.suffix.lower())
        sheet = sheets[0]
        from .dapodik import parse_sheet

        parsed = parse_sheet(sheet)
        if parsed.kind != "daftar_peserta_didik":
            return None
        stored = StoredFile(path=path, filename=path.name, size=path.stat().st_size, extension=path.suffix.lower())
        import_id = create_import(
            stored,
            sheet=sheet,
            parsed=parsed,
            fmt=format_label(path.suffix.lower()),
            actor="sistem",
            mode="upsert",
        )
        result = import_students(parsed, import_id=import_id, mode="upsert",
                                 actor="sistem", source_file=path.name)
        seed_ekskul_if_empty()
        log_audit("sistem", "admin", "impor_otomatis", "imports", import_id,
                  f"{result.imported} siswa dari {path.name}")
        return {"file": path.name, "imported": result.imported, "updated": result.updated,
                "import_id": import_id, "rows": len(parsed.records)}
    except Exception as exc:  # noqa: BLE001
        log.warning("Gagal impor otomatis berkas contoh: %s", exc)
        return None
