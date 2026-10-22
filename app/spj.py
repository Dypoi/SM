"""Pembuat berkas SPJ Tangerang Cerdas — biar tidak perlu menyalin satu per satu.

Menu TANGCER sudah bisa membaca SK penerima dan menyaring baris milik sekolah kita.
Modul ini melanjutkannya sampai **berkas SPJ siap tanda tangan**:

* ``PAKTA INTEGRITAS TANGERANG CERDAS.docx`` — isian nama, jabatan, sekolah, alamat, jumlah
  siswa, tahun, semester; tinggal ditempeli meterai;
* ``SURAT PERNYATAAN TRANSFER.docx`` — kop sekolah + tabel penerimaan/penyaluran/pengembalian;
* ``DAFTAR PENERIMA.xlsx`` (dan versi ``.pdf`` siap cetak) — NO · NOMOR REKENING · NAMA ·
  KELAS · L/P · NAMA ORANGTUA · NAMA SEKOLAH · BIAYA PERSONAL · VOLUME · JUMLAH ·
  TANDA TANGAN, sama seperti berkas kerja sekolah;
* lembar fotokopi buku tabungan + daftar sampel yang harus difotokopi.

Berkas DOCX dibuat sendiri dengan :mod:`zipfile` (tanpa pustaka tambahan) memakai
WordprocessingML seperlunya, dan bisa langsung dibuka di Word/LibreOffice.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import xml.sax.saxutils as saxutils
import zipfile
from dataclasses import dataclass, field as dataclass_field
from pathlib import Path
from typing import Any, Mapping, Sequence

from . import config, services, tangcer

#: Biaya personal Tangerang Cerdas untuk SMP/MTs per peserta didik per bulan.
BIAYA_PERSONAL_BAWAAN = 100_000
#: Volume bawaan: satu semester = 6 bulan.
VOLUME_BAWAAN = 6
#: Jumlah anak yang difotokopi buku tabungannya (mengikuti lembar sampel sekolah).
SAMPEL_BAWAAN = 10
#: Folder penyimpanan hasil SPJ di dalam folder data aplikasi.
FOLDER_SPJ = "spj"
#: Periode bulan tiap semester (dipakai di surat pernyataan transfer).
PERIODE_SEMESTER = {
    "1": "Januari s/d Juni",
    "2": "Juli s/d Desember",
    "ganjil": "Januari s/d Juni",
    "genap": "Juli s/d Desember",
}


def _belah(semester: str) -> tuple[str, str]:
    """Kembalikan ``(nomor semester, periode bulan)`` dari isian apa pun."""
    kunci = tangcer.norm_teks(semester).lower()
    nomor = "1" if kunci in ("1", "ganjil", "i", "satu") else "2" if kunci else ""
    if kunci in ("2", "genap", "ii", "dua"):
        nomor = "2"
    return nomor, PERIODE_SEMESTER.get(kunci, PERIODE_SEMESTER.get(nomor, "Januari s/d Juni"))


def samakan_semester(semester: Any) -> str:
    """Semester dipakai seragam sebagai «1» atau «2» (data TANGCER menyimpannya begitu)."""
    return _belah(str(semester or ""))[0]


# =========================================================================== #
# Isian & data
# =========================================================================== #
@dataclass
class DataSpj:
    """Semua isian yang dipakai menyusun berkas SPJ."""

    sekolah: str = ""
    npsn: str = ""
    alamat: str = ""
    kecamatan: str = ""
    kabupaten: str = ""
    provinsi: str = ""
    kepala: str = ""
    nip: str = ""
    jabatan: str = "Kepala Sekolah"
    tahun: str = ""
    semester: str = "1"
    tahap: str = "1"
    kota: str = ""
    tanggal: str = ""
    biaya_personal: int = BIAYA_PERSONAL_BAWAAN
    volume: int = VOLUME_BAWAAN
    sampel: int = SAMPEL_BAWAAN
    dana_masuk: bool = False
    catatan: list[str] = dataclass_field(default_factory=list)

    @property
    def periode(self) -> str:
        return _belah(self.semester)[1]

    @property
    def nomor_semester(self) -> str:
        return _belah(self.semester)[0] or self.semester

    @property
    def tanggal_teks(self) -> str:
        if self.tanggal:
            return self.tanggal
        return f"{self.kota}, {dt.datetime.now():%d %B %Y}".replace(
            "January", "Januari").replace("February", "Februari").replace(
            "March", "Maret").replace("April", "April").replace("May", "Mei").replace(
            "June", "Juni").replace("July", "Juli").replace("August", "Agustus").replace(
            "September", "September").replace("October", "Oktober").replace(
            "November", "November").replace("December", "Desember")

    @property
    def alamat_lengkap(self) -> str:
        bagian = [self.alamat, self.kecamatan, self.kabupaten, self.provinsi]
        return ", ".join(satu for satu in bagian if satu)


def _angka(nilai: Any, bawaan: int) -> int:
    teks = re.sub(r"[^0-9]", "", str(nilai or ""))
    return int(teks) if teks else bawaan


def data_dari(isian: Mapping[str, Any] | None = None) -> DataSpj:
    """Susun isian SPJ dari formulir, dilengkapi data sekolah yang tersimpan di aplikasi."""
    profil = services.school_profile()
    isian = dict(isian or {})
    data = DataSpj(
        sekolah=str(isian.get("sekolah") or profil["nama"] or "Sekolah"),
        npsn=str(isian.get("npsn") or profil["npsn"] or ""),
        alamat=str(isian.get("alamat") or profil["alamat"] or ""),
        kecamatan=str(isian.get("kecamatan") or profil["kecamatan"] or ""),
        kabupaten=str(isian.get("kabupaten") or profil["kabupaten"] or ""),
        provinsi=str(isian.get("provinsi") or profil["provinsi"] or ""),
        kepala=str(isian.get("kepala") or profil["kepala"] or "").strip(),
        nip=str(isian.get("nip") or "").strip(),
        jabatan=str(isian.get("jabatan") or "Kepala Sekolah").strip(),
        tahun=str(isian.get("tahun") or "").strip(),
        semester=samakan_semester(isian.get("semester") or profil["semester"]) or "1",
        tahap=str(isian.get("tahap") or "1").strip(),
        kota=str(isian.get("kota") or profil["kabupaten"] or "").strip(),
        tanggal=str(isian.get("tanggal") or "").strip(),
        biaya_personal=_angka(isian.get("biaya_personal"), BIAYA_PERSONAL_BAWAAN),
        volume=_angka(isian.get("volume"), VOLUME_BAWAAN),
        sampel=_angka(isian.get("sampel"), SAMPEL_BAWAAN),
        dana_masuk=bool(isian.get("dana_masuk")),
    )
    if not data.tahun:
        tahun_ajaran = str(profil.get("tahun_ajaran") or "")
        cocok = re.search(r"(20\d{2})", tahun_ajaran)
        data.tahun = cocok.group(1) if cocok else f"{dt.datetime.now():%Y}"
    if not data.kota:
        data.kota = "Tangerang"
    if not data.jabatan:
        data.jabatan = "Kepala Sekolah"
    if data.volume <= 0:
        data.volume = VOLUME_BAWAAN
    if data.biaya_personal <= 0:
        data.biaya_personal = BIAYA_PERSONAL_BAWAAN
    if data.sampel <= 0:
        data.sampel = SAMPEL_BAWAAN
    return data


def pilih_penerima(data: DataSpj, *, batas: int = 100_000) -> tuple[list[dict[str, Any]], list[str]]:
    """Ambil penerima tersimpan untuk tahun/semester ini, dengan catatan apa adanya.

    Isian tahun/semester waktu mengunggah berkas bisa berbeda dari isian SPJ (mis. «Ganjil»
    vs «1»), jadi pencariannya bertingkat dan tiap langkah dilaporkan — bukan diam-diam
    memakai data yang bukan maksud pengguna.
    """
    catatan: list[str] = []
    percobaan: list[tuple[dict[str, str], str]] = [
        ({"tahun": data.tahun, "semester": data.semester},
         f"tahun {data.tahun} semester {data.semester}"),
        ({"tahun": data.tahun}, f"tahun {data.tahun} (semua semester)"),
        ({}, "semua data penerima yang tersimpan"),
    ]
    for tapis, label in percobaan:
        hasil = tangcer.daftar_penerima(batas=batas, **tapis)
        if "semester" not in tapis and data.semester:
            hasil = [item for item in hasil
                     if samakan_semester(item.get("semester")) in ("", data.semester)]
        if hasil:
            catatan.append(f"{len(hasil)} penerima diambil dari data tersimpan ({label}).")
            if tapis.get("tahun") != data.tahun or tapis.get("semester") != data.semester:
                catatan.append("Tahun/semester pada isian SPJ belum sama dengan data yang "
                               "tersimpan — mohon diperiksa sebelum dicetak.")
            return hasil, catatan
    return [], ["Belum ada data penerima yang tersimpan di menu TANGCER."]


def baris_spj(data: DataSpj, *, batas: int = 100_000) -> list[dict[str, Any]]:
    """Baris SPJ: satu penerima tersimpan = satu baris siap cetak."""
    penerima, catatan = pilih_penerima(data, batas=batas)
    data.catatan = catatan
    keluaran: list[dict[str, Any]] = []
    for nomor, item in enumerate(penerima, start=1):
        orangtua = (item.get("nama_orangtua") or item.get("nama_ayah") or item.get("nama_ibu")
                    or "").strip()
        keluaran.append({
            "no": nomor,
            "id": item.get("id"),
            "rekening": (item.get("rekening") or "").strip(),
            "nama": (item.get("nama") or "").strip(),
            "rombel": (item.get("rombel") or "").strip(),
            "kelamin": (item.get("kelamin") or "").strip().upper()[:1],
            "orangtua": orangtua,
            "ayah": (item.get("nama_ayah") or "").strip(),
            "ibu": (item.get("nama_ibu") or "").strip(),
            "sekolah": (item.get("sekolah") or data.sekolah).strip(),
            "nisn": (item.get("nisn") or "").strip(),
            "biaya": data.biaya_personal,
            "volume": data.volume,
            "jumlah": data.biaya_personal * data.volume,
        })
    return keluaran


def ringkas_spj(data: DataSpj, baris: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    """Angka ringkasan yang dipakai di surat & di halaman pratinjau."""
    total = sum(int(item.get("jumlah") or 0) for item in baris)
    tanpa_rekening = sum(1 for item in baris if not item.get("rekening"))
    tanpa_orangtua = sum(1 for item in baris if not item.get("orangtua"))
    return {
        "jumlah_siswa": len(baris),
        "total": total,
        "tanpa_rekening": tanpa_rekening,
        "tanpa_orangtua": tanpa_orangtua,
        "laki_laki": sum(1 for item in baris if item.get("kelamin") == "L"),
        "perempuan": sum(1 for item in baris if item.get("kelamin") == "P"),
    }


# =========================================================================== #
# Penulis DOCX sederhana (WordprocessingML + zipfile, tanpa pustaka tambahan)
# =========================================================================== #
_NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MIME_DOKUMEN = ("application/vnd.openxmlformats-officedocument."
                 "wordprocessingml.document.main+xml")
_CONTENT_TYPES = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="{_MIME_DOKUMEN}"/>
<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
</Types>"""
_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
</Relationships>"""
_STYLES = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="{_NS_W}">
<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
<w:sz w:val="22"/><w:szCs w:val="22"/></w:rPr></w:rPrDefault>
<w:pPrDefault><w:pPr><w:spacing w:after="120"/></w:pPr></w:pPrDefault></w:docDefaults>
<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>
</w:styles>"""
#: Lebar halaman A4 & tepi 2,5 cm dalam satuan twips (1/20 pt).
_LEBAR_HALAMAN = 11906
_TINGGI_HALAMAN = 16838
_TEPI = 1418
_LEBAR_ISI = _LEBAR_HALAMAN - 2 * _TEPI


def _esc(teks: Any) -> str:
    return saxutils.escape("" if teks is None else str(teks))


def paragraf(teks: Any = "", *, tebal: bool = False, miring: bool = False, ukuran: int = 11,
             rata: str = "left", jarak: int = 120, garis_bawah: bool = False,
             masuk: int = 0) -> str:
    """Satu paragraf Word. ``rata``: left/center/right/justify."""
    gaya: list[str] = []
    if garis_bawah:
        gaya.append('<w:pBdr><w:bottom w:val="single" w:sz="12" w:space="2" '
                    'w:color="000000"/></w:pBdr>')
    if masuk:
        gaya.append(f'<w:ind w:left="{masuk}"/>')
    gaya.append(f'<w:spacing w:after="{jarak}"/>')
    gaya.append(f'<w:jc w:val="{rata}"/>')
    rpr = [f'<w:sz w:val="{ukuran * 2}"/>']
    if tebal:
        rpr.append("<w:b/>")
    if miring:
        rpr.append("<w:i/>")
    isi = ""
    if str(teks) != "":
        isi = (f'<w:r><w:rPr>{"".join(rpr)}</w:rPr>'
               f'<w:t xml:space="preserve">{_esc(teks)}</w:t></w:r>')
    return f'<w:p><w:pPr>{"".join(gaya)}</w:pPr>{isi}</w:p>'


def _sel(teks: Any, lebar: int, *, tebal: bool, rata: str) -> str:
    rpr = '<w:rPr><w:sz w:val="22"/><w:b/></w:rPr>' if tebal else '<w:rPr><w:sz w:val="22"/></w:rPr>'
    return (f'<w:tc><w:tcPr><w:tcW w:w="{lebar}" w:type="dxa"/></w:tcPr>'
            f'<w:p><w:pPr><w:spacing w:after="0"/><w:jc w:val="{rata}"/></w:pPr>'
            f'<w:r>{rpr}<w:t xml:space="preserve">{_esc(teks)}</w:t></w:r></w:p></w:tc>')


def tabel(lebar: Sequence[int], baris: Sequence[Sequence[Any]], *,
          tebal_judul: bool = True, rata: str = "left") -> str:
    """Tabel bergaris; ``lebar`` dalam twips dan sejajar dengan kolom."""
    grid = "".join(f'<w:gridCol w:w="{nilai}"/>' for nilai in lebar)
    potongan = [
        '<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:tblLayout w:type="fixed"/>'
        '<w:tblBorders>'
        '<w:top w:val="single" w:sz="6" w:color="000000"/>'
        '<w:left w:val="single" w:sz="6" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="6" w:color="000000"/>'
        '<w:right w:val="single" w:sz="6" w:color="000000"/>'
        '<w:insideH w:val="single" w:sz="6" w:color="000000"/>'
        '<w:insideV w:val="single" w:sz="6" w:color="000000"/>'
        '</w:tblBorders>'
        '<w:tblCellMar><w:top w:w="40" w:type="dxa"/><w:left w:w="80" w:type="dxa"/>'
        '<w:bottom w:w="40" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tblCellMar>'
        '</w:tblPr><w:tblGrid>%s</w:tblGrid>' % (sum(lebar), grid)
    ]
    for nomor, isi_baris in enumerate(baris):
        tebal = tebal_judul and nomor == 0
        sel = "".join(_sel(nilai, lebar[kolom] if kolom < len(lebar) else lebar[-1],
                           tebal=tebal, rata=rata if not tebal else "center")
                      for kolom, nilai in enumerate(isi_baris))
        potongan.append(f'<w:tr>{"".join(sel)}</w:tr>')
    potongan.append("</w:tbl>")
    return "".join(potongan)


def buat_docx(elemen: Sequence[str], *, judul: str = "Dokumen SPJ") -> bytes:
    """Rakit berkas ``.docx`` dari elemen paragraf/tabel (langsung bisa dibuka Word)."""
    inti = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:document xmlns:w="{_NS_W}"><w:body>{"".join(elemen)}'
            f'<w:sectPr><w:pgSz w:w="{_LEBAR_HALAMAN}" w:h="{_TINGGI_HALAMAN}"/>'
            f'<w:pgMar w:top="{_TEPI}" w:right="{_TEPI}" w:bottom="{_TEPI}" w:left="{_TEPI}" '
            f'w:header="720" w:footer="720" w:gutter="0"/></w:sectPr></w:body></w:document>')
    dibuat = f"{dt.datetime.now():%Y-%m-%dT%H:%M:%SZ}"
    core = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/'
            f'metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" '
            f'xmlns:dcterms="http://purl.org/dc/terms/" '
            f'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f'<dc:title>{_esc(judul)}</dc:title><dc:creator>Aplikasi SM</dc:creator>'
            f'<cp:lastModifiedBy>Aplikasi SM</cp:lastModifiedBy>'
            f'<dcterms:created xsi:type="dcterms:W3CDTF">{dibuat}</dcterms:created>'
            f'<dcterms:modified xsi:type="dcterms:W3CDTF">{dibuat}</dcterms:modified>'
            f'</cp:coreProperties>')
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as berkas:
        berkas.writestr("[Content_Types].xml", _CONTENT_TYPES)
        berkas.writestr("_rels/.rels", _RELS)
        berkas.writestr("word/document.xml", inti)
        berkas.writestr("word/styles.xml", _STYLES)
        berkas.writestr("docProps/core.xml", core)
    return buffer.getvalue()


# =========================================================================== #
# Dokumen-dokumen SPJ
# =========================================================================== #
def _kepala_teks(data: DataSpj) -> str:
    return data.kepala or "………………………………"


def _nip_teks(data: DataSpj) -> str:
    return f"NIP. {data.nip}" if data.nip else "NIP. …………………."


def dokumen_pakta(data: DataSpj, ringkas: Mapping[str, Any]) -> bytes:
    """Pakta Integritas (surat pernyataan tanggung jawab mutlak) terisi otomatis."""
    elemen = [
        paragraf("PAKTA INTEGRITAS", tebal=True, ukuran=14, rata="center", jarak=40),
        paragraf("(SURAT PERNYATAAN TANGGUNG JAWAB MUTLAK)", rata="center", jarak=240),
        paragraf("Yang bertanda tangan dibawah ini :", jarak=80),
        paragraf(f"Nama : {_kepala_teks(data)}", jarak=40),
        paragraf(f"Jabatan : {data.jabatan}", jarak=40),
        paragraf(f"Nama Sekolah : {data.sekolah}", jarak=40),
        paragraf(f"Alamat : {data.alamat_lengkap or '………………………………'}", jarak=240),
        paragraf("Dengan ini menyatakan bahwa :", jarak=80),
        paragraf(f"1. Jumlah data Penerima Tangerang Cerdas Tahun {data.tahun} "
                  f"Semester {data.nomor_semester} Tahap {data.tahap} sebagaimana terlampir "
                  f"adalah benar siswa kami sejumlah {ringkas.get('jumlah_siswa', 0)} Siswa;",
                  jarak=80),
        paragraf(f"2. Penerima bantuan merupakan siswa dari keluarga tidak mampu yang "
                  f"dibuktikan dari data DSEN Tahun {data.tahun};", jarak=80),
        paragraf("3. Saya bertanggung jawab mutlak secara hukum atas kebenaran dan keabsahan "
                 "dokumen administrasi data yang diajukan.", jarak=240),
        paragraf("Demikian pernyataan ini saya buat dengan sebenarnya, dalam keadaan sehat "
                 "jasmani/rohani dan tanpa paksaan dari siapapun.", jarak=400),
        paragraf(data.tanggal_teks, rata="right", jarak=40),
        paragraf(f"Kepala {data.sekolah}" if data.sekolah else "Kepala Sekolah",
                  rata="right", jarak=40),
        paragraf("", jarak=240),
        paragraf("Materai 10.000", rata="right", jarak=240),
        paragraf(_kepala_teks(data), tebal=True, rata="right", jarak=20),
        paragraf(_nip_teks(data), rata="right", jarak=40),
    ]
    return buat_docx(elemen, judul=f"Pakta Integritas {data.sekolah}")


def _rupiah(nilai: Any) -> str:
    try:
        angka = float(nilai or 0)
    except (TypeError, ValueError):
        angka = 0.0
    return f"{angka:,.0f}".replace(",", ".")


def dokumen_pernyataan(data: DataSpj, ringkas: Mapping[str, Any]) -> bytes:
    """Surat Pernyataan Transfer (kop sekolah + tabel penerimaan/penyaluran)."""
    total = int(ringkas.get("total") or 0)
    titik = "……………………….."
    if data.dana_masuk:
        penerimaan, pengembalian = _rupiah(total), _rupiah(0)
    else:
        penerimaan, pengembalian = titik, titik
    baris_tabel = [
        ["No.", "Waktu", "Penerimaan (Rp.)", "Penyaluran (Rp.)", "Pengembalian (Rp.)"],
        ["1.", f"Semester {data.nomor_semester} (Periode {data.periode}) {data.tahun}",
         penerimaan, _rupiah(total), pengembalian],
        ["", "Jumlah", penerimaan, _rupiah(total), pengembalian],
    ]
    lebar = [500, 2600, 1900, 1900, 2170]
    catatan = [
        paragraf("SURAT PERNYATAAN TRANSFER", tebal=True, ukuran=14, rata="center", jarak=240),
        paragraf("Yang bertanda tangan dibawah ini :", jarak=80),
        paragraf(f"Nama : {_kepala_teks(data)}", jarak=40),
        paragraf(f"NIP : {data.nip or '…………………………'}", jarak=40),
        paragraf(f"Jabatan : {data.jabatan}", jarak=40),
        paragraf(f"Alamat : {data.alamat_lengkap or '………………………………'}", jarak=240),
        paragraf("Dengan ini menyatakan bahwa :", jarak=80),
        paragraf("Belanja pembiayaan Tangerang Cerdas telah ditransfer dan telah dikonfirmasi "
                 f"kepada Peserta Didik, yang masih aktif dan berstatus sebagai Peserta Didik "
                 f"pada {data.sekolah}, dan tidak ada pemotongan dengan dalih apapun kepada "
                 "Peserta Didik.", rata="justify", jarak=120),
        paragraf("Kami akan memantau penyaluran/pentransferan Biaya Tangerang Cerdas kepada "
                 "Peserta Didik setelah uang masuk kedalam rekening Peserta Didik, dan "
                 "mengembalikan ketika kelebihan dan/tidak terserap. Kami akan menyerahkan "
                 "pelaporan/SPJ Tangerang Cerdas setelah Dinas Pendidikan melaksanakan "
                 "Pentransferan Dana Kepada Peserta Didik. Penerima dan Penyaluran Belanja "
                 "Pembiayaan Tangerang Cerdas di sekolah kami adalah sebagai berikut :",
                 rata="justify", jarak=200),
        tabel(lebar, baris_tabel),
        paragraf("", jarak=120),
        paragraf("Kami bersedia menyalurkan Besaran Pembiayaan Tangerang Cerdas sesuai dengan "
                 "petunjuk teknis yaitu :", jarak=80),
        tabel([2400, 6670], [["SMP/MTs/SKh", f"Personal : {_rupiah(data.biaya_personal)}/"
                                              f"peserta didik/bulan"]]),
        paragraf("", jarak=120),
        paragraf("Apabila dikemudian hari pernyataan ini tidak sesuai dengan keadaan yang "
                 "sebenarnya, saya bersedia dikenakan sanksi administrasi dan/atau dituntut "
                 "ganti rugi dan/atau tuntutan lainnya sesuai dengan ketentuan dan peraturan "
                 "perundang-undangan yang berlaku.", rata="justify", jarak=120),
        paragraf("Demikian surat pernyataan ini dibuat dengan sebenarnya untuk dipergunakan "
                 "sebagaimana mestinya.", rata="justify", jarak=400),
        paragraf(data.tanggal_teks, rata="right", jarak=40),
        paragraf("Kepala Sekolah", rata="right", jarak=40),
        paragraf(data.sekolah, rata="right", jarak=240),
        paragraf("Materai 10.000", rata="right", jarak=240),
        paragraf(_kepala_teks(data), tebal=True, rata="right", jarak=20),
        paragraf(_nip_teks(data), rata="right", jarak=40),
    ]
    # Kop sekolah di atas judul (diberi garis pemisah seperti kop surat).
    kop = [paragraf(data.sekolah, tebal=True, ukuran=14, rata="center", jarak=20),
           paragraf(data.alamat_lengkap, ukuran=10, rata="center", jarak=20, garis_bawah=True),
           paragraf("", jarak=120)]
    return buat_docx(kop + catatan, judul=f"Surat Pernyataan Transfer {data.sekolah}")


def dokumen_fotokopi(data: DataSpj, baris: Sequence[Mapping[str, Any]]) -> bytes:
    """Lembar kosong untuk fotokopi buku tabungan & foto dokumentasi (mengikuti contoh)."""
    elemen = [
        paragraf("FOTOCOPY IDENTITAS BUKU TABUNGAN PENERIMA TANGERANG CERDAS",
                 tebal=True, ukuran=12, rata="center", jarak=240),
    ]
    elemen.extend(paragraf("", jarak=240) for _ in range(6))
    elemen.append(paragraf("FOTO DOKUMENTASI SAAT TANDA TANGAN PENERIMA TANGERANG CERDAS",
                           tebal=True, ukuran=12, rata="center", jarak=240))
    elemen.extend(paragraf("", jarak=240) for _ in range(2))
    elemen.append(paragraf(f"Jumlah penerima pada daftar SPJ: {len(baris)} siswa "
                           f"({data.sekolah}, Semester {data.nomor_semester} {data.tahun}).",
                           ukuran=10, miring=True, jarak=40))
    return buat_docx(elemen, judul="Fotokopi Buku Tabungan")


def dokumen_sampel(data: DataSpj, baris: Sequence[Mapping[str, Any]]) -> bytes:
    """Daftar anak yang difotokopi buku tabungannya (sampel), siap dicentang."""
    sampel = list(baris)[:max(1, data.sampel)]
    isi = [["No.", "Nama Siswa", "Kelas", "Nama Orang Tua", "Nomor Rekening", "Centang"]]
    for nomor, item in enumerate(sampel, start=1):
        isi.append([nomor, item.get("nama") or "", item.get("rombel") or "",
                    item.get("orangtua") or "", item.get("rekening") or "", "☐"])
    lebar = [500, 2500, 700, 2200, 2100, 1070]
    elemen = [
        paragraf("DAFTAR SAMPEL DOKUMENTASI & FOTOKOPI BUKU TABUNGAN", tebal=True, ukuran=13,
                 rata="center", jarak=40),
        paragraf(f"{data.sekolah} · Semester {data.nomor_semester} ({data.periode}) {data.tahun} "
                 f"· {len(sampel)} dari {len(baris)} siswa", rata="center", jarak=200),
        tabel(lebar, isi),
        paragraf("", jarak=120),
        paragraf("Nama yang tercantum di atas difotokopi buku tabungannya (halaman identitas) "
                 "dan difoto saat menandatangani tanda terima.", miring=True, jarak=40),
    ]
    return buat_docx(elemen, judul="Daftar Sampel Dokumentasi")


def xlsx_daftar(data: DataSpj, baris: Sequence[Mapping[str, Any]]) -> bytes:
    """Daftar penerima dalam Excel — bentuknya mengikuti berkas «DAFTAR PENERIMA.xlsx»."""
    import openpyxl  # noqa: PLC0415
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    buku = openpyxl.Workbook()
    lembar = buku.active
    lembar.title = "Daftar Penerima"
    judul = ("DAFTAR PENERIMA PEMBAYARAN TANGERANG CERDAS JENJANG SMP/MTs/SKh "
             "SWASTA DAN NEGERI")
    sub = (f"SEMESTER {data.nomor_semester} ( {data.periode.upper()} ) "
           f"TAHUN ANGGARAN {data.tahun}")
    lembar["A1"] = judul
    lembar["A1"].font = Font(bold=True, size=12)
    lembar["A1"].alignment = Alignment(horizontal="center")
    lembar.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
    lembar["A2"] = sub
    lembar["A2"].font = Font(bold=True, size=11)
    lembar["A2"].alignment = Alignment(horizontal="center")
    lembar.merge_cells(start_row=2, start_column=1, end_row=2, end_column=11)

    kepala = ["NO", "NOMOR REKENING", "NAMA", "KELAS", "L/P", "NAMA ORANGTUA",
              "NAMA SEKOLAH", "BIAYA PERSONAL", "VOLUME", "JUMLAH", "TANDA TANGAN"]
    isi_kepala = ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10 = 8x9", "11"]
    tepi = Side(style="thin", color="000000")
    for kolom, label in enumerate(kepala, start=1):
        sel = lembar.cell(row=4, column=kolom, value=label)
        sel.font = Font(bold=True)
        sel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sel.fill = PatternFill("solid", fgColor="E8EEF9")
        sel.border = Border(left=tepi, right=tepi, top=tepi, bottom=tepi)
        sel2 = lembar.cell(row=5, column=kolom, value=isi_kepala[kolom - 1])
        sel2.font = Font(bold=True, size=9)
        sel2.alignment = Alignment(horizontal="center")
        sel2.border = Border(left=tepi, right=tepi, top=tepi, bottom=tepi)
    lembar.merge_cells(start_row=4, start_column=11, end_row=4, end_column=11)

    mulai = 6
    for nomor, item in enumerate(baris, start=mulai):
        nilai = [nomor - mulai + 1, item.get("rekening") or "", item.get("nama") or "",
                 item.get("rombel") or "", item.get("kelamin") or "",
                 item.get("orangtua") or "", item.get("sekolah") or data.sekolah,
                 int(item.get("biaya") or data.biaya_personal),
                 int(item.get("volume") or data.volume), f"=H{nomor}*I{nomor}", ""]
        for kolom, angka in enumerate(nilai, start=1):
            sel = lembar.cell(row=nomor, column=kolom, value=angka)
            sel.border = Border(left=tepi, right=tepi, top=tepi, bottom=tepi)
            sel.alignment = Alignment(horizontal="left" if kolom in (2, 3, 6, 7) else "center",
                                      vertical="center", wrap_text=kolom in (3, 6, 7))
            if kolom in (8, 10):
                sel.number_format = "#,##0"
    baris_akhir = max(mulai, mulai + len(baris) - 1)
    kaki = baris_akhir + 2
    lembar.cell(row=kaki, column=9, value=data.tanggal_teks).alignment = Alignment(
        horizontal="left")
    lembar.cell(row=kaki + 1, column=9, value="Kepala Sekolah")
    lembar.cell(row=kaki + 4, column=9, value=_kepala_teks(data)).font = Font(bold=True)
    lembar.cell(row=kaki + 5, column=9, value=_nip_teks(data))
    for kolom, lebar in {"A": 4.5, "B": 19.8, "C": 32.8, "D": 9.2, "E": 6.0, "F": 20.2,
                         "G": 22.0, "H": 14.7, "I": 9.3, "J": 12.0, "K": 13.2}.items():
        lembar.column_dimensions[kolom].width = lebar
    lembar.freeze_panes = "A6"
    lembar.print_title_rows = "4:5"
    buffer = io.BytesIO()
    buku.save(buffer)
    return buffer.getvalue()


def pdf_daftar(data: DataSpj, baris: Sequence[Mapping[str, Any]]) -> bytes:
    """Daftar penerima siap cetak (PDF) — untuk ditandatangani tanpa mengubah Excel."""
    from .pdf import buat_pdf_tabel  # noqa: PLC0415

    kolom = [("No", 26), ("Nomor Rekening", 110), ("Nama Penerima", 132), ("Kelas", 36),
             ("L/P", 26), ("Nama Orang Tua", 108), ("Jumlah (Rp)", 85)]
    isi = [[item.get("no"), item.get("rekening") or "-", item.get("nama") or "",
            item.get("rombel") or "", item.get("kelamin") or "",
            item.get("orangtua") or "-", _rupiah(item.get("jumlah"))] for item in baris]
    ringkas = ringkas_spj(data, baris)
    return buat_pdf_tabel(
        "DAFTAR PENERIMA TANGERANG CERDAS",
        kolom,
        isi,
        subjudul=[data.sekolah,
                  f"Semester {data.nomor_semester} ({data.periode}) Tahun {data.tahun} · "
                  f"{ringkas['jumlah_siswa']} siswa · total Rp {_rupiah(ringkas['total'])}",
                  f"Biaya personal Rp {_rupiah(data.biaya_personal)} per bulan × "
                  f"{data.volume} bulan"],
        catatan_kaki=[f"Kepala {data.sekolah} — {_kepala_teks(data)} ({_nip_teks(data)})",
                      f"Dicetak dari aplikasi SM · {dt.datetime.now():%d-%m-%Y %H:%M}"],
    )


# =========================================================================== #
# Paket berkas & penyimpanan
# =========================================================================== #
def _folder_spj() -> Path:
    folder = config.DATA_DIR / FOLDER_SPJ
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _nama_aman(nama: str) -> str:
    return re.sub(r"[^A-Za-z0-9._ ()-]+", "-", (nama or "berkas").strip())[:90] or "berkas"


#: Nama berkas paket SPJ (urutannya = urutan tampil di halaman).
URUTAN_BERKAS = (
    "PAKTA INTEGRITAS TANGERANG CERDAS.docx",
    "SURAT PERNYATAAN TRANSFER.docx",
    "DAFTAR PENERIMA.xlsx",
    "DAFTAR PENERIMA.pdf",
    "FOTOCOPY BUKU TABUNGAN dan Dokumentasi Tanda Terima ( SAMPEL 10 ANAK ).docx",
    "DAFTAR SAMPEL DOKUMENTASI.docx",
)
#: Keterangan singkat tiap berkas untuk halaman hasil.
LABEL_BERKAS = {
    URUTAN_BERKAS[0]: "Nama, jabatan, sekolah, alamat, jumlah siswa sudah terisi — tinggal "
                      "ditempeli meterai 10.000 dan ditandatangani.",
    URUTAN_BERKAS[1]: "Kop sekolah + tabel penerimaan/penyaluran/pengembalian sudah terisi "
                      "angka dari daftar penerima.",
    URUTAN_BERKAS[2]: "Excel 11 kolom seperti berkas kerja sekolah; kolom JUMLAH memakai "
                      "rumus =biaya×volume.",
    URUTAN_BERKAS[3]: "Versi siap cetak untuk ditandatangani (ukuran A4, berkepala).",
    URUTAN_BERKAS[4]: "Lembar kosong untuk ditempeli fotokopi buku tabungan & foto dokumentasi.",
    URUTAN_BERKAS[5]: "Daftar anak yang difotokopi (sampel) — biar tidak ada yang tertukar.",
}


def buat_paket(data: DataSpj, baris: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    """Buat seluruh berkas SPJ, simpan di folder data, kembalikan daftar unduhannya."""
    if not baris:
        raise tangcer.TangcerError(
            f"Belum ada penerima tersimpan untuk tahun {data.tahun} semester "
            f"{data.semester}. Unggah & simpan daftar penerima TANGCER dulu, lalu buat SPJ."
        )
    ringkas = ringkas_spj(data, baris)
    isi_berkas = {
        URUTAN_BERKAS[0]: dokumen_pakta(data, ringkas),
        URUTAN_BERKAS[1]: dokumen_pernyataan(data, ringkas),
        URUTAN_BERKAS[2]: xlsx_daftar(data, baris),
        URUTAN_BERKAS[3]: pdf_daftar(data, baris),
        URUTAN_BERKAS[4]: dokumen_fotokopi(data, baris),
        URUTAN_BERKAS[5]: dokumen_sampel(data, baris),
    }
    token = f"spj-{dt.datetime.now():%Y%m%d-%H%M%S}"
    folder = _folder_spj() / token
    folder.mkdir(parents=True, exist_ok=True)
    keluaran: list[dict[str, Any]] = []
    for nama in URUTAN_BERKAS:
        isi = isi_berkas[nama]
        aman = _nama_aman(nama)
        (folder / aman).write_bytes(isi)
        keluaran.append({"nama": nama, "berkas": aman, "ukuran": len(isi),
                         "label": LABEL_BERKAS.get(nama, "")})
    return {"token": token, "folder": str(folder), "sekolah": data.sekolah,
            "tahun": data.tahun, "semester": data.semester, "ringkas": ringkas,
            "baris": list(baris), "berkas": keluaran, "catatan": list(data.catatan),
            "dibuat": dt.datetime.now().strftime("%d-%m-%Y %H:%M")}


def berkas_spj(token: str, nama: str) -> Path | None:
    """Cari satu berkas hasil SPJ (menolak nama yang keluar dari foldernya)."""
    if not token or not nama or "/" in nama or "\\" in nama or nama.startswith("."):
        return None
    if "/" in token or "\\" in token or token.startswith("."):
        return None
    path = _folder_spj() / token / nama
    return path if path.is_file() else None


def nama_berkas_paket(token: str) -> dict[str, Any] | None:
    """Keterangan paket yang pernah dibuat (dibaca ulang dari isi foldernya)."""
    for paket in paket_tersimpan(batas=200):
        if paket["token"] == token:
            return paket
    return None


def hapus_paket(token: str) -> bool:
    """Hapus satu folder paket SPJ (dipakai tombol bersihkan di halaman)."""
    if not token or "/" in token or "\\" in token or token.startswith("."):
        return False
    folder = _folder_spj() / token
    if not folder.is_dir():
        return False
    for berkas in folder.iterdir():
        if berkas.is_file():
            berkas.unlink(missing_ok=True)
    folder.rmdir()
    return True


def paket_tersimpan(batas: int = 10) -> list[dict[str, Any]]:
    """Daftar paket SPJ yang pernah dibuat (terbaru dulu) — untuk halaman TANGCER."""
    folder = _folder_spj()
    keluaran: list[dict[str, Any]] = []
    for satu in sorted(folder.iterdir(), reverse=True)[:batas] if folder.is_dir() else []:
        if not satu.is_dir():
            continue
        berkas = [{"nama": item.name, "ukuran": item.stat().st_size} for item in
                  sorted(satu.iterdir()) if item.is_file()]
        keluaran.append({"token": satu.name,
                         "dibuat": dt.datetime.fromtimestamp(satu.stat().st_mtime)
                         .strftime("%d-%m-%Y %H:%M"),
                         "berkas": berkas})
    return keluaran
