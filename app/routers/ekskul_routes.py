"""Modul ekstrakurikuler."""

from __future__ import annotations

import csv
import datetime as dt
import io
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .. import config

from fastapi import APIRouter, Depends, Form, Request
from fastapi.responses import RedirectResponse, Response

from .. import auth, services
from ..pdf import buat_pdf_tabel
from ..web import render

router = APIRouter()

#: Jabatan yang biasa dipakai dalam ekstrakurikuler.
JABATAN_OPTIONS = ("Anggota", "Ketua", "Wakil Ketua", "Sekretaris", "Bendahara", "Pelatih")


def _redirect(pesan: str, level: str = "ok", target: str = "/ekstrakurikuler", fragmen: str = ""):
    """Alihkan halaman disertai pesan.

    ``target`` boleh sudah memuat query (mis. hasil "Cari cepat siswa"), jadi
    pemisah & ditambahkan, dan ``fragmen`` menaruh penanda bagian halaman.
    """
    pemisah = "&" if "?" in target else "?"
    ujung = f"#{fragmen}" if fragmen else ""
    return RedirectResponse(f"{target}{pemisah}level={level}&msg={quote_plus(pesan)}{ujung}",
                            status_code=303)


def _boleh_kelola(user: auth.SessionUser, ekskul_id: int) -> bool:
    """Petugas sekolah bebas; akun ekskul hanya untuk ekskulnya sendiri."""
    if user.is_staff:
        return True
    return bool(user.is_ekskul and user.ekskul_id == ekskul_id)


def _tolak_kelola(user: auth.SessionUser):
    """Akun ekskul yang membuka ekskul lain diarahkan ke ekskulnya sendiri."""
    return _redirect(
        "Akun Anda hanya dapat mengelola ekstrakurikuler sendiri.",
        level="warn",
        target=user.halaman_ekskul,
    )


@router.get("/ekstrakurikuler")
def daftar_ekskul(request: Request, user: auth.SessionUser = Depends(auth.require_staff)):
    edit_id = request.query_params.get("edit", "")
    edit_item = services.get_ekskul(int(edit_id)) if edit_id.isdigit() else None
    return render(
        request,
        "ekskul/list.html",
        {
            "page_title": "Ekstrakurikuler",
            "daftar": services.list_ekskul(),
            "stats": services.ekskul_stats(),
            "hari_options": services.HARI_OPTIONS,
            "edit_item": edit_item,
            "ringkasan": services.ekskul_ringkas(limit=8),
            "jumlah_akun": {baris["id"]: services.akun_ekskul_ekskul(int(baris["id"]))
                            for baris in services.list_ekskul()},
        },
    )


@router.post("/ekstrakurikuler/simpan")
def simpan_ekskul(
    request: Request,
    ekskul_id: str = Form(""),
    nama: str = Form(""),
    pembina: str = Form(""),
    pelatih: str = Form(""),
    hari: str = Form(""),
    jam_mulai: str = Form(""),
    jam_selesai: str = Form(""),
    deskripsi: str = Form(""),
    aktif: str = Form("0"),
    user: auth.SessionUser = Depends(auth.require_staff),
):
    data = {
        "nama": nama.strip(),
        "pembina": pembina.strip() or None,
        "pelatih": pelatih.strip() or None,
        "hari": hari or None,
        "jam_mulai": jam_mulai or None,
        "jam_selesai": jam_selesai or None,
        "deskripsi": deskripsi.strip() or None,
        "aktif": 1 if aktif in {"1", "on", "true"} else 0,
    }
    try:
        new_id = services.save_ekskul(data, ekskul_id=int(ekskul_id) if ekskul_id.isdigit() else None,
                                      actor=user.username)
    except ValueError as exc:
        return _redirect(str(exc), level="err")

    if ekskul_id.isdigit():
        return _redirect(f"Ekstrakurikuler {data['nama']} diperbarui.", target=f"/ekstrakurikuler/{new_id}")
    return _redirect(f"Ekstrakurikuler {data['nama']} ditambahkan.", target=f"/ekstrakurikuler/{new_id}")


@router.get("/ekstrakurikuler/{ekskul_id}")
def detail_ekskul(request: Request, ekskul_id: int, rombel: str = "", cari: str = "",
                  user: auth.SessionUser = Depends(auth.require_user)):
    ekskul = services.get_ekskul(ekskul_id)
    if ekskul is None:
        return render(request, "error.html", {"kode": 404, "pesan": "Ekstrakurikuler tidak ditemukan."}, status_code=404)
    if not _boleh_kelola(user, ekskul_id):
        return _tolak_kelola(user)
    rombel = (rombel or "").strip()
    cari = (cari or "").strip()
    ada_cari = bool(rombel or cari)
    return render(
        request,
        "ekskul/detail.html",
        {
            "page_title": ekskul["nama"],
            "ekskul": ekskul,
            "anggota": services.ekskul_members(ekskul_id),
            "hari_options": services.HARI_OPTIONS,
            "opsi_rombel": services.distinct_values("rombel"),
            "akun_ekskul": services.akun_ekskul_ekskul(ekskul_id),
            "boleh_ubah": user.is_staff,
            "pilihan_siswa": services.pilihan_siswa_ekskul(),
            "nilai_options": services.EKSKUL_NILAI,
            "jabatan_options": JABATAN_OPTIONS,
            "rombel_dipilih": rombel,
            "kata_cari": cari,
            "cari_aktif": ada_cari,
            "hasil_cari": services.cari_siswa_cepat(ekskul_id, rombel=rombel, cari=cari) if ada_cari else [],
        },
    )


@router.post("/ekstrakurikuler/{ekskul_id}/hapus")
def hapus_ekskul(request: Request, ekskul_id: int, user: auth.SessionUser = Depends(auth.require_staff)):
    ekskul = services.get_ekskul(ekskul_id)
    services.delete_ekskul(ekskul_id, actor=user.username)
    return _redirect(f"Ekstrakurikuler {ekskul['nama'] if ekskul else ''} dihapus.")


@router.post("/ekstrakurikuler/{ekskul_id}/anggota")
def tambah_anggota(
    request: Request,
    ekskul_id: int,
    student_id: str = Form(""),
    nisn: str = Form(""),
    jabatan: str = Form("Anggota"),
    kembali_rombel: str = Form(""),
    kembali_cari: str = Form(""),
    user: auth.SessionUser = Depends(auth.require_user),
):
    if not _boleh_kelola(user, ekskul_id):
        return _tolak_kelola(user)
    target = f"/ekstrakurikuler/{ekskul_id}"
    # Bila penambahan dilakukan dari hasil "Cari cepat siswa", kembalikan ke hasil
    # pencarian itu supaya pembina dapat memasukkan beberapa siswa sekaligus.
    rombel_kembali = (kembali_rombel or "").strip()[:20]
    cari_kembali = (kembali_cari or "").strip()[:60]
    fragmen = ""
    if rombel_kembali or cari_kembali:
        target += f"?rombel={quote_plus(rombel_kembali)}&cari={quote_plus(cari_kembali)}"
        fragmen = "cari-cepat"
    siswa = None
    if student_id.isdigit():
        siswa = services.get_student(int(student_id))
    if siswa is None and nisn.strip():
        # Menerima NISN maupun nama siswa (nama yang kembar harus memakai NISN).
        siswa, galat = services.cari_siswa_ekskul(nisn)

    if siswa is None:
        return _redirect(galat or "Siswa tidak ditemukan. Masukkan NISN yang benar.", level="err",
                         target=target, fragmen=fragmen)

    member_id = services.add_ekskul_member(ekskul_id, siswa["id"], jabatan=jabatan or "Anggota",
                                           actor=user.username)
    if member_id is None:
        return _redirect(f"{siswa['nama']} sudah terdaftar di ekstrakurikuler ini.", level="warn",
                         target=target, fragmen=fragmen)
    return _redirect(f"{siswa['nama']} ({siswa.get('rombel') or '-'}) masuk ke ekstrakurikuler.",
                     target=target, fragmen=fragmen)


@router.post("/ekstrakurikuler/anggota/{member_id}/hapus")
def hapus_anggota(request: Request, member_id: int,
                  user: auth.SessionUser = Depends(auth.require_user)):
    row = services.db.query_one("SELECT ekskul_id FROM ekskul_members WHERE id = ?", (member_id,))
    if row and not _boleh_kelola(user, int(row["ekskul_id"])):
        return _tolak_kelola(user)
    services.remove_ekskul_member(member_id, actor=user.username)
    target = f"/ekstrakurikuler/{row['ekskul_id']}" if row else "/ekstrakurikuler"
    return _redirect("Anggota dikeluarkan dari ekstrakurikuler.", target=target)


@router.post("/ekstrakurikuler/anggota/{member_id}")
def perbarui_anggota(
    request: Request,
    member_id: int,
    jabatan: str = Form("Anggota"),
    nilai: str = Form(""),
    catatan: str = Form(""),
    status: str = Form("aktif"),
    user: auth.SessionUser = Depends(auth.require_user),
):
    row = services.db.query_one("SELECT ekskul_id FROM ekskul_members WHERE id = ?", (member_id,))
    if row and not _boleh_kelola(user, int(row["ekskul_id"])):
        return _tolak_kelola(user)
    nilai_bersih = (nilai or "").strip().upper()[:1]
    data = {
        "jabatan": jabatan or "Anggota",
        # Nilai ekstrakurikuler memakai huruf A-D (sesuai permintaan sekolah);
        # kolom angka lama tetap dibiarkan kosong agar tidak menyesatkan.
        "predikat": nilai_bersih if nilai_bersih in services.EKSKUL_NILAI else None,
        "catatan": (catatan or "").strip() or None,
        "status": status or "aktif",
    }
    services.update_ekskul_member(member_id, data, actor=user.username)
    target = f"/ekstrakurikuler/{row['ekskul_id']}" if row else "/ekstrakurikuler"
    return _redirect("Data anggota diperbarui.", target=target)


#: Kolom daftar anggota untuk semua bentuk ekspor (CSV/Excel/PDF).
KOLOM_ANGGOTA: tuple[tuple[str, int], ...] = (
    ("No", 28), ("Nama", 140), ("NISN", 66), ("Rombel", 40),
    ("JK", 26), ("Jabatan", 62), ("Nilai", 30), ("Catatan", 131),
)


def _data_anggota(ekskul_id: int) -> tuple[dict, list[list[object]]]:
    """(ekskul, baris) untuk keperluan ekspor."""
    ekskul = services.get_ekskul(ekskul_id) or {}
    baris: list[list[object]] = []
    for index, member in enumerate(services.ekskul_members(ekskul_id), start=1):
        baris.append([
            index, member["nama"], member["nisn"] or "", member["rombel"] or "",
            member["jk"] or "", member["jabatan"] or "Anggota",
            member.get("predikat") or "", member.get("catatan") or "",
        ])
    return ekskul, baris


def _subjudul_ekspor(ekskul: dict) -> list[str]:
    profil = services.school_profile()
    jadwal = ""
    if ekskul.get("hari"):
        jadwal = f" · {ekskul['hari']}"
        if ekskul.get("jam_mulai"):
            jadwal += f" {ekskul['jam_mulai']}"
            if ekskul.get("jam_selesai"):
                jadwal += f"-{ekskul['jam_selesai']}"
    baris = [
        str(profil.get("nama") or ""),
        f"Ekstrakurikuler {ekskul.get('nama') or ''}{jadwal}",
    ]
    pendamping = []
    if ekskul.get("pembina"):
        pendamping.append(f"Pembina: {ekskul['pembina']}")
    if ekskul.get("pelatih"):
        pendamping.append(f"Pelatih: {ekskul['pelatih']}")
    if pendamping:
        baris.append(" · ".join(pendamping))
    baris.append(f"Tahun ajaran {services.school_profile().get('tahun_ajaran', '-')}"
                 f" · semester {services.school_profile().get('semester', '-')}")
    return baris


def _berkas_aman(teks: str) -> str:
    return "".join(karakter if karakter.isalnum() else "-" for karakter in (teks or "ekskul")).strip("-")


@router.get("/ekstrakurikuler/{ekskul_id}/anggota.csv")
def ekspor_anggota_csv(request: Request, ekskul_id: int,
                       user: auth.SessionUser = Depends(auth.require_user)):
    ekskul = services.get_ekskul(ekskul_id)
    if ekskul is None:
        return RedirectResponse("/ekstrakurikuler", status_code=303)
    if not _boleh_kelola(user, ekskul_id):
        return _tolak_kelola(user)
    _, baris = _data_anggota(ekskul_id)
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([label for label, _ in KOLOM_ANGGOTA])
    writer.writerows(baris)
    nama_berkas = f"anggota-{_berkas_aman(ekskul['nama'])}.csv"
    return Response(
        content=buffer.getvalue().encode("utf-8-sig"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{nama_berkas}"'},
    )


@router.get("/ekstrakurikuler/{ekskul_id}/anggota.xlsx")
def ekspor_anggota_xlsx(request: Request, ekskul_id: int,
                        user: auth.SessionUser = Depends(auth.require_user)):
    ekskul = services.get_ekskul(ekskul_id)
    if ekskul is None:
        return RedirectResponse("/ekstrakurikuler", status_code=303)
    if not _boleh_kelola(user, ekskul_id):
        return _tolak_kelola(user)
    _, baris = _data_anggota(ekskul_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Anggota"
    ws.append([f"Daftar Anggota Ekstrakurikuler {ekskul['nama']}"])
    ws["A1"].font = Font(bold=True, size=13)
    for baris_sub in _subjudul_ekspor(ekskul)[1:]:
        ws.append([baris_sub])
    ws.append([])
    ws.append([label for label, _ in KOLOM_ANGGOTA])
    for sel in ws[ws.max_row]:
        sel.font = Font(bold=True)
        sel.alignment = Alignment(horizontal="center")
    for item in baris:
        ws.append(item)
    for urutan, (label, lebar) in enumerate(KOLOM_ANGGOTA, start=1):
        ws.column_dimensions[ws.cell(row=1, column=urutan).column_letter].width = max(9, lebar // 7)
    ws.freeze_panes = ws.cell(row=ws.max_row - len(baris) + 1, column=1)

    buffer = io.BytesIO()
    wb.save(buffer)
    nama_berkas = f"anggota-{_berkas_aman(ekskul['nama'])}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nama_berkas}"'},
    )


@router.get("/ekstrakurikuler/{ekskul_id}/anggota.pdf")
def ekspor_anggota_pdf(request: Request, ekskul_id: int,
                       user: auth.SessionUser = Depends(auth.require_user)):
    ekskul = services.get_ekskul(ekskul_id)
    if ekskul is None:
        return RedirectResponse("/ekstrakurikuler", status_code=303)
    if not _boleh_kelola(user, ekskul_id):
        return _tolak_kelola(user)
    _, baris = _data_anggota(ekskul_id)
    pdf = buat_pdf_tabel(
        judul=f"Daftar Anggota Ekstrakurikuler {ekskul['nama']}",
        subjudul=_subjudul_ekspor(ekskul),
        kolom=[(label, float(lebar)) for label, lebar in KOLOM_ANGGOTA],
        baris=baris,
        catatan_kaki=[
            f"{config.APP_NAME} · {services.school_profile().get('nama', '')}",
            f"Dicetak {dt.datetime.now().strftime('%d-%m-%Y %H:%M')} · "
            "Tanda tangan pembina/pelatih: ____________________",
        ],
    )
    nama_berkas = f"anggota-{_berkas_aman(ekskul['nama'])}.pdf"
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{nama_berkas}"'},
    )
