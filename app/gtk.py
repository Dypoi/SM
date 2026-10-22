"""Impor data **Guru & Tenaga Kependidikan (GTK)** dari berkas mentah Dapodik.

Berkas «daftar guru» / «daftar tendik» yang diunduh dari Dapodik bentuknya sama
seperti berkas peserta didik:

* baris 1-4  : judul, nama sekolah, wilayah, tanggal unduh & pengunduh
* baris 5    : header kolom (51 kolom: No, Nama, NUPTK, JK, …)
* baris 6+   : data guru / tendik

Alurnya sama seperti impor lain di aplikasi ini: **unggah → pratinjau → simpan**,
supaya tidak ada data yang masuk sebelum dilihat manusia. Nilai dibaca **apa
adanya** mengikuti isi berkas; yang dinormalisasi hanya format tanggal
(→ ISO ``YYYY-MM-DD``) dan jawaban Ya/Tidak.

Kunci pengenal tiap baris (supaya unggahan ulang tidak menggandakan data):
``NUPTK`` → bila kosong ``NIP`` → bila kosong ``NIK`` → bila semuanya kosong,
gabungan nama + tanggal lahir (dengan peringatan).
"""

from __future__ import annotations

import io
import json
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Sequence

from . import db
from .dapodik import (
    FieldSpec,
    Issue,
    convert_value,
    extract_metadata,
    normalize_header,
)
from .readers import (
    SheetData,
    UnsupportedFormatError,
    clean_text,
    format_label,
    read_spreadsheet,
    supported_extensions,
)

log = logging.getLogger("sm.gtk")


class GtkError(RuntimeError):
    """Berkas GTK tidak bisa dibaca / tidak dikenali."""


# --------------------------------------------------------------------------- #
# Definisi field GTK (mengikuti 51 kolom ekspor Dapodik)
# --------------------------------------------------------------------------- #
GTK_FIELDS: tuple[FieldSpec, ...] = (
    # ---- Identitas ----
    FieldSpec("nama", "Nama", ("nama", "nama lengkap", "nama ptk", "nama guru", "nama tendik"), "identitas", required=True),
    FieldSpec("nuptk", "NUPTK", ("nuptk", "nomor unik pendidik"), "identitas"),
    FieldSpec("jk", "JK", ("jk", "jenis kelamin", "l p", "lp"), "identitas"),
    FieldSpec("tempat_lahir", "Tempat Lahir", ("tempat lahir",), "identitas"),
    FieldSpec("tanggal_lahir", "Tanggal Lahir", ("tanggal lahir", "tgl lahir"), "identitas", kind="date"),
    FieldSpec("nip", "NIP", ("nip", "nomor induk pegawai"), "identitas"),
    FieldSpec("agama", "Agama", ("agama",), "identitas"),
    FieldSpec("nik", "NIK", ("nik", "nik pegawai"), "identitas"),
    FieldSpec("no_kk", "No KK", ("no kk", "nomor kk", "no kartu keluarga", "nomor kartu keluarga"), "identitas"),
    FieldSpec("kewarganegaraan", "Kewarganegaraan", ("kewarganegaraan", "warga negara", "kewarganegaraan wni"), "identitas"),

    # ---- Kepegawaian ----
    FieldSpec("status_kepegawaian", "Status Kepegawaian", ("status kepegawaian", "status pegawai"), "kepegawaian"),
    FieldSpec("jenis_ptk", "Jenis PTK", ("jenis ptk", "jenis pegawai"), "kepegawaian"),
    FieldSpec("tugas_tambahan", "Tugas Tambahan", ("tugas tambahan",), "kepegawaian"),
    FieldSpec("sk_cpns", "SK CPNS", ("sk cpns", "nomor sk cpns"), "kepegawaian"),
    FieldSpec("tanggal_cpns", "Tanggal CPNS", ("tanggal cpns", "tgl cpns"), "kepegawaian", kind="date"),
    FieldSpec("sk_pengangkatan", "SK Pengangkatan", ("sk pengangkatan", "nomor sk pengangkatan"), "kepegawaian"),
    FieldSpec("tmt_pengangkatan", "TMT Pengangkatan", ("tmt pengangkatan",), "kepegawaian", kind="date"),
    FieldSpec("lembaga_pengangkatan", "Lembaga Pengangkatan", ("lembaga pengangkatan",), "kepegawaian"),
    FieldSpec("pangkat_golongan", "Pangkat Golongan", ("pangkat golongan", "pangkat", "golongan"), "kepegawaian"),
    FieldSpec("sumber_gaji", "Sumber Gaji", ("sumber gaji",), "kepegawaian"),
    FieldSpec("tmt_pns", "TMT PNS", ("tmt pns",), "kepegawaian", kind="date"),
    FieldSpec("lisensi_kepala_sekolah", "Sudah Lisensi Kepala Sekolah", ("sudah lisensi kepala sekolah", "lisensi kepala sekolah"), "kepegawaian", kind="ya_tidak"),
    FieldSpec("pernah_diklat", "Pernah Diklat Kepengawasan", ("pernah diklat kepengawasan", "diklat kepengawasan"), "kepegawaian", kind="ya_tidak"),
    FieldSpec("braille", "Keahlian Braille", ("keahlian braille", "braille"), "kepegawaian", kind="ya_tidak"),
    FieldSpec("isyarat", "Keahlian Bahasa Isyarat", ("keahlian bahasa isyarat", "bahasa isyarat"), "kepegawaian", kind="ya_tidak"),
    FieldSpec("nuks", "NUKS", ("nuks",), "kepegawaian"),

    # ---- Alamat & kontak ----
    FieldSpec("alamat", "Alamat Jalan", ("alamat jalan", "alamat", "jalan"), "alamat"),
    FieldSpec("rt", "RT", ("rt",), "alamat"),
    FieldSpec("rw", "RW", ("rw",), "alamat"),
    FieldSpec("dusun", "Nama Dusun", ("nama dusun", "dusun"), "alamat"),
    FieldSpec("kelurahan", "Desa/Kelurahan", ("desa kelurahan", "kelurahan", "desa"), "alamat"),
    FieldSpec("kecamatan", "Kecamatan", ("kecamatan",), "alamat"),
    FieldSpec("kode_pos", "Kode Pos", ("kode pos", "kodepos"), "alamat"),
    FieldSpec("telepon", "Telepon", ("telepon", "telp", "no telepon"), "alamat"),
    FieldSpec("hp", "HP", ("hp", "no hp", "handphone", "ponsel"), "alamat"),
    FieldSpec("email", "Email", ("email", "e mail", "surel"), "alamat"),
    FieldSpec("lintang", "Lintang", ("lintang", "latitude"), "alamat", kind="float"),
    FieldSpec("bujur", "Bujur", ("bujur", "longitude"), "alamat", kind="float"),

    # ---- Keluarga ----
    FieldSpec("ibu_kandung", "Nama Ibu Kandung", ("nama ibu kandung", "ibu kandung"), "keluarga"),
    FieldSpec("status_perkawinan", "Status Perkawinan", ("status perkawinan", "status kawin"), "keluarga"),
    FieldSpec("pasangan_nama", "Nama Suami/Istri", ("nama suami istri", "nama pasangan", "nama suami", "nama istri"), "keluarga"),
    FieldSpec("pasangan_nip", "NIP Suami/Istri", ("nip suami istri", "nip pasangan"), "keluarga"),
    FieldSpec("pasangan_pekerjaan", "Pekerjaan Suami/Istri", ("pekerjaan suami istri", "pekerjaan pasangan"), "keluarga"),

    # ---- Pajak & bank ----
    FieldSpec("npwp", "NPWP", ("npwp",), "keuangan"),
    FieldSpec("nama_wp", "Nama Wajib Pajak", ("nama wajib pajak",), "keuangan"),
    FieldSpec("bank", "Bank", ("bank", "nama bank"), "keuangan"),
    FieldSpec("rekening", "Nomor Rekening Bank", ("nomor rekening bank", "nomor rekening", "no rekening"), "keuangan"),
    FieldSpec("rekening_nama", "Rekening Atas Nama", ("rekening atas nama",), "keuangan"),
    FieldSpec("karpeg", "Karpeg", ("karpeg",), "keuangan"),
    FieldSpec("karis_karsu", "Karis/Karsu", ("karis karsu", "karis", "karsu"), "keuangan"),
)

GTK_BY_KEY: dict[str, FieldSpec] = {spec.key: spec for spec in GTK_FIELDS}

GROUP_LABELS = {
    "identitas": "Identitas",
    "kepegawaian": "Kepegawaian",
    "alamat": "Alamat & Kontak",
    "keluarga": "Keluarga",
    "keuangan": "Pajak & Bank",
    "lainnya": "Lainnya",
}

GROUP_ORDER = ["identitas", "kepegawaian", "alamat", "keluarga", "keuangan"]


def fields_by_group() -> dict[str, list[FieldSpec]]:
    grouped: dict[str, list[FieldSpec]] = {key: [] for key in GROUP_ORDER}
    for spec in GTK_FIELDS:
        grouped.setdefault(spec.group, []).append(spec)
    return {key: value for key, value in grouped.items() if value}


def _build_alias_index() -> dict[str, list[FieldSpec]]:
    index: dict[str, list[FieldSpec]] = {}
    for spec in GTK_FIELDS:
        for alias in (spec.label, *spec.aliases):
            index.setdefault(normalize_header(alias), []).append(spec)
    return index


_ALIAS_INDEX = _build_alias_index()


# --------------------------------------------------------------------------- #
# Hasil parsing
# --------------------------------------------------------------------------- #
@dataclass
class GtkRecord:
    row_number: int
    values: dict[str, Any] = field(default_factory=dict)
    raw: dict[str, Any] = field(default_factory=dict)
    issues: list[Issue] = field(default_factory=list)

    @property
    def has_error(self) -> bool:
        return any(issue.level == "error" for issue in self.issues)


@dataclass
class HasilGtk:
    sheet_name: str
    kind: str                       # 'daftar_gtk' | 'generik'
    header_row: int
    data_start_row: int
    jenis_berkas: str               # 'guru' | 'tendik' | 'campuran' | ''
    columns: list[str]
    mapping: dict[str, str]         # field_key -> label kolom asli
    column_map: dict[int, str]      # index kolom -> field_key
    records: list[GtkRecord]
    issues: list[Issue]
    metadata: dict[str, str]
    total_columns: int = 0

    @property
    def mapped_field_count(self) -> int:
        return len(self.column_map)

    @property
    def detected_fields(self) -> list[str]:
        return list(self.mapping.keys())


# --------------------------------------------------------------------------- #
# Deteksi struktur berkas
# --------------------------------------------------------------------------- #
def _skor_header(baris: Sequence[Any]) -> int:
    skor = 0
    for sel in baris:
        normal = normalize_header(sel)
        if normal and normal in _ALIAS_INDEX:
            skor += 1
    return skor


def find_header_row(sheet: SheetData, scan_limit: int = 40) -> int:
    """Cari baris header (1-based) pada lembar GTK."""
    terbaik, skor_terbaik = 0, 0
    batas = min(sheet.row_count, scan_limit)
    for index in range(batas):
        skor = _skor_header(sheet.rows[index])
        if skor > skor_terbaik:
            terbaik, skor_terbaik = index, skor
    if skor_terbaik >= 3:
        return terbaik + 1
    # Tidak ada header GTK: pakai baris pertama yang cukup terisi.
    for index in range(batas):
        terisi = sum(1 for sel in sheet.rows[index] if sel is not None and str(sel).strip() != "")
        if terisi >= 2:
            return index + 1
    return 1


def build_mapping(columns: Sequence[str]) -> dict[int, str]:
    """Petakan indeks kolom -> field_key (tanpa duplikasi field)."""
    column_map: dict[int, str] = {}
    terpakai: set[str] = set()
    for index, label in enumerate(columns):
        for spec in _candidates_for(label):
            if spec.key in terpakai:
                continue
            column_map[index] = spec.key
            terpakai.add(spec.key)
            break
    return column_map


def _candidates_for(label: str) -> list[FieldSpec]:
    normal = normalize_header(label)
    if not normal:
        return []
    tepat = list(_ALIAS_INDEX.get(normal, []))
    if tepat:
        tepat.sort(key=lambda spec: GTK_FIELDS.index(spec))
        return tepat
    # Pencocokan longgar: label yang memuat alias (mis. «NUPTK/NIP»).
    longgar: list[FieldSpec] = []
    for alias, specs in _ALIAS_INDEX.items():
        if len(alias) >= 5 and alias in normal:
            longgar.extend(specs)
    longgar.sort(key=lambda spec: GTK_FIELDS.index(spec))
    return longgar


def detect_kind(mapping: dict[int, str]) -> str:
    kunci = set(mapping.values())
    if "nama" in kunci and (kunci & {"nuptk", "nip", "nik", "jenis_ptk", "status_kepegawaian"}):
        return "daftar_gtk"
    return "generik"


def _jenis_dari_judul(metadata: dict[str, str], sheet_name: str) -> str:
    teks = f"{metadata.get('judul', '')} {sheet_name}".lower()
    if "tenaga kependidikan" in teks or "tendik" in teks:
        return "tendik"
    if "guru" in teks:
        return "guru"
    return ""


def kunci_ptk(values: dict[str, Any]) -> str:
    """Kunci pengenal baris: NUPTK → NIP → NIK → nama+tanggal lahir."""
    for kunci in ("nuptk", "nip", "nik"):
        nilai = values.get(kunci)
        if nilai is not None and str(nilai).strip():
            return str(nilai).strip()
    nama = str(values.get("nama") or "").strip().upper()
    lahir = str(values.get("tanggal_lahir") or "").strip()
    return f"nama:{nama}|lahir:{lahir}"


# --------------------------------------------------------------------------- #
# Validasi
# --------------------------------------------------------------------------- #
def validate_record(record: GtkRecord, terlihat: dict[str, int]) -> None:
    values = record.values
    nama = values.get("nama")
    if not nama or len(str(nama)) < 3:
        record.issues.append(Issue(record.row_number, "error", "nama",
                                   "Nama kosong atau terlalu pendek."))
    if not any(str(values.get(kunci) or "").strip() for kunci in ("nuptk", "nip", "nik")):
        record.issues.append(Issue(record.row_number, "warning", "nuptk",
                                   "Tanpa NUPTK/NIP/NIK — baris ini dikenali lewat nama "
                                   "sehingga rawan ganda bila ada nama sama."))
    elif not values.get("nuptk"):
        record.issues.append(Issue(record.row_number, "info", "nuptk", "NUPTK belum diisi."))
    jk = values.get("jk")
    if jk and str(jk).strip().upper() not in {"L", "P"}:
        record.issues.append(Issue(record.row_number, "warning", "jk",
                                   f"Jenis kelamin '{jk}' di luar L/P."))
    nik = values.get("nik")
    if nik and str(nik).strip().isdigit() and len(str(nik).strip()) != 16:
        record.issues.append(Issue(record.row_number, "warning", "nik",
                                   f"NIK {nik} berjumlah {len(str(nik).strip())} digit "
                                   "(seharusnya 16)."))
    kunci = kunci_ptk(values)
    if kunci in terlihat and terlihat[kunci] != record.row_number:
        record.issues.append(Issue(record.row_number, "warning", None,
                                   f"Kunci pengenal ganda dengan baris {terlihat[kunci]} "
                                   "(NUPTK/NIP/NIK sama)."))


# --------------------------------------------------------------------------- #
# API utama parser
# --------------------------------------------------------------------------- #
def parse_sheet(sheet: SheetData, *, header_row: int | None = None) -> HasilGtk:
    """Urai satu worksheet menjadi catatan GTK + temuan validasi."""
    kepala = header_row or find_header_row(sheet)
    atas = sheet.rows[kepala - 1] if kepala - 1 < sheet.row_count else []
    columns = [clean_text(sel) or f"Kolom {i + 1}" for i, sel in enumerate(atas)]
    mapping = build_mapping(columns)
    data_start = kepala + 1

    records: list[GtkRecord] = []
    issues: list[Issue] = []
    terlihat: dict[str, int] = {}

    for index in range(data_start - 1, sheet.row_count):
        baris = sheet.rows[index]
        nomor = index + 1
        if all(sel is None or str(sel).strip() == "" for sel in baris):
            continue
        values: dict[str, Any] = {}
        raw: dict[str, Any] = {}
        temuan: list[Issue] = []
        for posisi, sel in enumerate(baris):
            label = columns[posisi] if posisi < len(columns) else f"Kolom {posisi + 1}"
            if sel is None or str(sel).strip() == "":
                continue
            raw[label] = sel
            kunci = mapping.get(posisi)
            if not kunci:
                continue
            spec = GTK_BY_KEY.get(kunci)
            if spec is None:
                continue
            ubah, peringatan = convert_value(spec, sel)
            if ubah is not None:
                values[kunci] = ubah
            if peringatan:
                temuan.append(Issue(nomor, "warning", kunci, peringatan))
        if not values:
            if raw:
                temuan.append(Issue(nomor, "info", None,
                                    "Baris tidak menghasilkan data yang dikenal (dilewati)."))
                issues.extend(temuan)
            continue
        if values.get("nama"):
            values["nama"] = clean_text(values["nama"], upper=True)
        kode = kunci_ptk(values)
        terlihat.setdefault(kode, nomor)
        record = GtkRecord(row_number=nomor, values=values, raw=raw, issues=temuan)
        validate_record(record, terlihat)
        records.append(record)

    for record in records:
        issues.extend(record.issues)

    metadata = extract_metadata(sheet, kepala)
    jenis = _jenis_dari_judul(metadata, sheet.name)
    if not jenis and records:
        macam = {str(item.values.get("jenis_ptk") or "").strip().lower() for item in records}
        macam.discard("")
        if macam == {"guru"}:
            jenis = "guru"
        elif macam and macam <= {"tenaga kependidikan", "kepala sekolah"}:
            jenis = "tendik"
        elif macam:
            jenis = "campuran"

    detail = {key: columns[pos] for pos, key in sorted(mapping.items()) if pos < len(columns)}
    return HasilGtk(
        sheet_name=sheet.name,
        kind=detect_kind(mapping),
        header_row=kepala,
        data_start_row=data_start,
        jenis_berkas=jenis,
        columns=columns,
        mapping=detail,
        column_map=dict(mapping),
        records=records,
        issues=issues,
        metadata=metadata,
        total_columns=len(columns),
    )


def ringkas_temuan(temuan: Sequence[Issue]) -> dict[str, int]:
    hasil = {"error": 0, "warning": 0, "info": 0}
    for item in temuan:
        hasil[item.level] = hasil.get(item.level, 0) + 1
    return hasil


# --------------------------------------------------------------------------- #
# Pembacaan berkas
# --------------------------------------------------------------------------- #
EKSTENSI_GTK = tuple(supported_extensions())


def baca_berkas(path: str | Path, ext: str | None = None) -> tuple[SheetData, HasilGtk, str]:
    """Baca berkas spreadsheet lalu urai lembar pertamanya sebagai data GTK."""
    berkas = Path(path)
    ekstensi = (ext or berkas.suffix.lower()).lower()
    if ekstensi not in EKSTENSI_GTK:
        raise UnsupportedFormatError(
            f"Jenis berkas «{ekstensi or 'tanpa ekstensi'}» belum didukung menu ini. "
            "Yang didukung: " + ", ".join(EKSTENSI_GTK)
        )
    try:
        lembar = read_spreadsheet(berkas, ekstensi)
    except (UnsupportedFormatError, GtkError):
        raise
    except Exception as galat:  # noqa: BLE001 - pesan asli diteruskan ke UI
        raise GtkError(f"Gagal membaca berkas: {galat}") from galat
    sheet = lembar[0]
    return sheet, parse_sheet(sheet), format_label(ekstensi)


def baca_isi(nama: str, isi: bytes) -> tuple[HasilGtk, str, str]:
    """Varian ``baca_berkas`` untuk isi berkas yang sudah ada di memori."""
    from . import services

    simpan = services.save_upload(nama, isi)
    try:
        sheet, hasil, label = baca_berkas(simpan.path, simpan.extension)
    except Exception:
        simpan.path.unlink(missing_ok=True)
        raise
    return hasil, label, str(simpan.path)


# --------------------------------------------------------------------------- #
# Penyimpanan
# --------------------------------------------------------------------------- #
KOLOM_TULIS: tuple[str, ...] = tuple(spec.key for spec in GTK_FIELDS)


def _catatan_baca(sheet: SheetData, hasil: HasilGtk, label: str) -> list[str]:
    ringkas = ringkas_temuan(hasil.issues)
    baris = [
        f"Dibaca sebagai {label}: lembar «{sheet.name}», "
        f"judul kolom pada baris {hasil.header_row}, data mulai baris {hasil.data_start_row}.",
        f"{len(hasil.records)} baris data, {hasil.mapped_field_count} kolom dikenali "
        f"dari {hasil.total_columns} kolom berkas.",
    ]
    if hasil.metadata.get("sekolah"):
        baris.append(f"Sekolah pada kepala berkas: {hasil.metadata['sekolah']}.")
    if hasil.metadata.get("tanggal_unduh"):
        baris.append(f"Tanggal unduh pada berkas: {hasil.metadata['tanggal_unduh']}.")
    if hasil.jenis_berkas:
        nama_jenis = {"guru": "daftar guru", "tendik": "daftar tendik",
                      "campuran": "campuran guru & tendik"}[hasil.jenis_berkas]
        baris.append(f"Berkas dikenali sebagai {nama_jenis}.")
    if hasil.kind != "daftar_gtk":
        baris.append("PERHATIAN — struktur kolomnya tidak mirip daftar GTK Dapodik "
                     "(kolom Nama + NUPTK/NIP/NIK tidak ditemukan); periksa pratinjau "
                     "baik-baik sebelum menyimpan.")
    tak_dikenal = [label for pos, label in enumerate(hasil.columns)
                   if label and not label.startswith("Kolom ") and pos not in hasil.column_map]
    if tak_dikenal:
        contoh = ", ".join(f"«{nama}»" for nama in tak_dikenal[:6])
        sisa = f" (+{len(tak_dikenal) - 6} lagi)" if len(tak_dikenal) > 6 else ""
        baris.append(f"Kolom yang tidak dikenali (tidak disimpan): {contoh}{sisa}.")
    if ringkas["error"] or ringkas["warning"]:
        baris.append(f"Temuan pemeriksaan: {ringkas['error']} galat, "
                     f"{ringkas['warning']} peringatan.")
    return baris


def buat_impor(*, nama_berkas: str, path: Path | None, ext: str, ukuran: int,
               sheet: SheetData, hasil: HasilGtk, fmt: str,
               actor: str | None) -> int:
    return db.insert_returning_id(
        """
        INSERT INTO gtk_imports(
            berkas, ext, ukuran, sumber_path, sekolah, judul, tanggal_unduh,
            jenis_berkas, sheet, baris_judul, total_baris, kolom, catatan,
            status, actor
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            nama_berkas, ext, ukuran, str(path) if path else None,
            hasil.metadata.get("sekolah", ""), hasil.metadata.get("judul", ""),
            hasil.metadata.get("tanggal_unduh", ""), hasil.jenis_berkas,
            sheet.name, hasil.header_row, len(hasil.records),
            json.dumps(hasil.mapping, ensure_ascii=False),
            "\n".join(_catatan_baca(sheet, hasil, fmt)),
            "pratinjau", actor,
        ),
    )


def get_impor(import_id: int) -> dict[str, Any] | None:
    baris = db.query_one("SELECT * FROM gtk_imports WHERE id = ?", (import_id,))
    hasil = db.row_to_dict(baris)
    if hasil is not None and hasil.get("kolom"):
        try:
            hasil["kolom"] = json.loads(hasil["kolom"])
        except (ValueError, TypeError):
            hasil["kolom"] = {}
    return hasil


def daftar_impor(limit: int = 25) -> list[dict[str, Any]]:
    return db.rows_to_dicts(
        db.query_all("SELECT * FROM gtk_imports ORDER BY id DESC LIMIT ?", (limit,))
    )


def tandai_impor(import_id: int, *, status: str, baru: int | None = None,
                 diperbarui: int | None = None, dilewati: int | None = None,
                 gagal: int | None = None, catatan: str | None = None) -> None:
    ubah: dict[str, Any] = {"status": status}
    if baru is not None:
        ubah["disimpan_baru"] = baru
    if diperbarui is not None:
        ubah["diperbarui"] = diperbarui
    if dilewati is not None:
        ubah["dilewati"] = dilewati
    if gagal is not None:
        ubah["gagal"] = gagal
    if catatan is not None:
        ubah["catatan"] = catatan
    if status == "tersimpan":
        klausa = ", ".join(f"{kunci} = ?" for kunci in ubah)
        db.execute(f"UPDATE gtk_imports SET {klausa}, disimpan = datetime('now','localtime') "
                   "WHERE id = ?", (*ubah.values(), import_id))
    else:
        klausa = ", ".join(f"{kunci} = ?" for kunci in ubah)
        db.execute(f"UPDATE gtk_imports SET {klausa} WHERE id = ?", (*ubah.values(), import_id))


def hapus_impor(import_id: int) -> None:
    kepala = get_impor(import_id)
    if kepala and kepala.get("sumber_path"):
        try:
            Path(kepala["sumber_path"]).unlink(missing_ok=True)
        except OSError:  # pragma: no cover
            pass
    db.execute("DELETE FROM gtk_pegawai WHERE import_id = ?", (import_id,))
    db.execute("DELETE FROM gtk_imports WHERE id = ?", (import_id,))


def _sama(a: Any, b: Any) -> bool:
    if a is None and b is None:
        return True
    if isinstance(a, str) and isinstance(b, str):
        return a.strip().casefold() == b.strip().casefold()
    if a is None or b is None:
        return False
    try:
        return float(a) == float(b)
    except (TypeError, ValueError):
        return str(a).strip() == str(b).strip()


def simpan_banyak(import_id: int, daftar: Sequence[GtkRecord], *,
                  source_file: str | None = None,
                  only_valid: bool = True) -> dict[str, Any]:
    """Simpan catatan GTK (tambah baru / perbarui yang kuncinya sama)."""
    hasil = {"baru": 0, "diperbarui": 0, "dilewati": 0, "gagal": 0, "daftar_gagal": []}
    with db.transaction() as conn:
        for record in daftar:
            if only_valid and record.has_error:
                hasil["gagal"] += 1
                pesan = record.issues[0].message if record.issues else "data tidak valid"
                hasil["daftar_gagal"].append(f"Baris {record.row_number}: {pesan}")
                continue
            values = {kunci: nilai for kunci, nilai in record.values.items()
                      if kunci in KOLOM_TULIS}
            if not values.get("nama"):
                hasil["gagal"] += 1
                hasil["daftar_gagal"].append(f"Baris {record.row_number}: tanpa nama.")
                continue
            kunci = kunci_ptk(values)
            ada = conn.execute("SELECT id FROM gtk_pegawai WHERE kunci = ?", (kunci,)).fetchone()
            if ada is None:
                kolom = ["kunci", *values.keys(), "source_file", "import_id"]
                tanda = ", ".join("?" for _ in kolom)
                conn.execute(
                    f"INSERT INTO gtk_pegawai({', '.join(kolom)}) VALUES({tanda})",
                    [kunci, *values.values(), source_file, import_id],
                )
                hasil["baru"] += 1
                continue
            kini = dict(conn.execute("SELECT * FROM gtk_pegawai WHERE id = ?",
                                     (ada["id"],)).fetchone())
            ubah = {kunci_kolom: nilai for kunci_kolom, nilai in values.items()
                    if not _sama(kini.get(kunci_kolom), nilai)}
            if not ubah:
                hasil["dilewati"] += 1
                continue
            ubah["source_file"] = source_file
            ubah["import_id"] = import_id
            klausa = ", ".join(f"{nama} = ?" for nama in ubah)
            conn.execute(f"UPDATE gtk_pegawai SET {klausa}, diubah = datetime('now','localtime') "
                         "WHERE id = ?", (*ubah.values(), ada["id"])),
            hasil["diperbarui"] += 1
    hasil["daftar_gagal"] = hasil["daftar_gagal"][:10]
    return hasil


def hapus_satu(pegawai_id: int) -> None:
    db.execute("DELETE FROM gtk_pegawai WHERE id = ?", (pegawai_id,))


def hapus_pegawai(*, jenis: str = "", status: str = "", cari: str = "") -> int:
    klausa, nilai = _klausa_filter(jenis=jenis, status=status, cari=cari)
    cur = db.execute(f"DELETE FROM gtk_pegawai{klausa}", nilai)
    return int(cur.rowcount or 0)


def get_pegawai(pegawai_id: int) -> dict[str, Any] | None:
    return db.row_to_dict(db.query_one("SELECT * FROM gtk_pegawai WHERE id = ?", (pegawai_id,)))


def _klausa_filter(*, jenis: str = "", status: str = "", cari: str = "") -> tuple[str, list[Any]]:
    syarat: list[str] = []
    nilai: list[Any] = []
    if jenis:
        syarat.append("jenis_ptk = ?")
        nilai.append(jenis)
    if status:
        syarat.append("status_kepegawaian = ?")
        nilai.append(status)
    cari = cari.strip()
    if cari:
        syarat.append("(nama LIKE ? OR nuptk LIKE ? OR nip LIKE ? OR nik LIKE ? OR hp LIKE ?)")
        nilai.extend([f"%{cari}%"] * 5)
    return (" WHERE " + " AND ".join(syarat)) if syarat else "", nilai


def daftar_pegawai(*, jenis: str = "", status: str = "", cari: str = "",
                   batas: int = 500) -> list[dict[str, Any]]:
    klausa, nilai = _klausa_filter(jenis=jenis, status=status, cari=cari)
    return db.rows_to_dicts(db.query_all(
        f"SELECT * FROM gtk_pegawai{klausa} ORDER BY nama COLLATE NOCASE LIMIT ?",
        (*nilai, batas),
    ))


def daftar_jenis() -> list[str]:
    baris = db.query_all(
        "SELECT DISTINCT jenis_ptk FROM gtk_pegawai WHERE jenis_ptk IS NOT NULL "
        "AND TRIM(jenis_ptk) <> '' ORDER BY jenis_ptk")
    return [row[0] for row in baris]


def daftar_status() -> list[str]:
    baris = db.query_all(
        "SELECT DISTINCT status_kepegawaian FROM gtk_pegawai WHERE status_kepegawaian IS NOT NULL "
        "AND TRIM(status_kepegawaian) <> '' ORDER BY status_kepegawaian")
    return [row[0] for row in baris]


def ringkas_pegawai() -> dict[str, Any]:
    total = int(db.query_value("SELECT COUNT(*) FROM gtk_pegawai", default=0) or 0)

    def hitung(kolom: str, nilai_kolom: str) -> int:
        return int(db.query_value(f"SELECT COUNT(*) FROM gtk_pegawai WHERE {kolom} = ?",
                                  (nilai_kolom,), default=0) or 0)

    per_jenis = db.rows_to_dicts(db.query_all(
        "SELECT jenis_ptk AS jenis, COUNT(*) AS jumlah FROM gtk_pegawai "
        "GROUP BY jenis_ptk ORDER BY jumlah DESC"))
    per_status = db.rows_to_dicts(db.query_all(
        "SELECT status_kepegawaian AS status, COUNT(*) AS jumlah FROM gtk_pegawai "
        "GROUP BY status_kepegawaian ORDER BY jumlah DESC"))
    return {
        "total": total,
        "guru": hitung("jenis_ptk", "Guru"),
        "tendik": hitung("jenis_ptk", "Tenaga Kependidikan"),
        "kepala_sekolah": hitung("jenis_ptk", "Kepala Sekolah"),
        "tanpa_nuptk": int(db.query_value(
            "SELECT COUNT(*) FROM gtk_pegawai WHERE nuptk IS NULL OR TRIM(nuptk) = ''",
            default=0) or 0),
        "per_jenis": per_jenis,
        "per_status": per_status,
    }


# --------------------------------------------------------------------------- #
# Ekspor
# --------------------------------------------------------------------------- #
KOLOM_EKSPOR: tuple[tuple[str, str], ...] = (
    ("no", "No"),
    *((spec.key, spec.label) for spec in GTK_FIELDS),
)


def baris_ekspor(*, jenis: str = "", status: str = "", cari: str = "") -> list[list[Any]]:
    daftar = daftar_pegawai(jenis=jenis, status=status, cari=cari, batas=100_000)
    hasil: list[list[Any]] = []
    for nomor, item in enumerate(daftar, start=1):
        baris: list[Any] = [nomor]
        for spec in GTK_FIELDS:
            nilai = item.get(spec.key)
            baris.append("" if nilai is None else nilai)
        hasil.append(baris)
    return hasil


def ekspor_csv(*, jenis: str = "", status: str = "", cari: str = "") -> bytes:
    import csv

    buffer = io.StringIO()
    tulis = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    tulis.writerow([label for _, label in KOLOM_EKSPOR])
    tulis.writerows(baris_ekspor(jenis=jenis, status=status, cari=cari))
    return buffer.getvalue().encode("utf-8-sig")


def ekspor_xlsx(*, jenis: str = "", status: str = "", cari: str = "") -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buku = openpyxl.Workbook()
    sheet = buku.active
    sheet.title = "Guru & Tendik"
    sheet.sheet_view.showGridLines = False

    kepala_font = Font(bold=True, color="FFFFFF", size=10)
    kepala_isi = PatternFill("solid", fgColor="1F6F43")
    for posisi, (_, label) in enumerate(KOLOM_EKSPOR, start=1):
        sel = sheet.cell(row=1, column=posisi, value=label)
        sel.font = kepala_font
        sel.fill = kepala_isi
        sel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for baris in baris_ekspor(jenis=jenis, status=status, cari=cari):
        sheet.append(baris)
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(KOLOM_EKSPOR))}{sheet.max_row}"
    lebar = {"No": 5, "Nama": 30, "NUPTK": 19, "NIP": 20, "NIK": 19, "No KK": 19,
             "Alamat Jalan": 34, "Email": 30, "SK Pengangkatan": 20, "NPWP": 20}
    for posisi, (_, label) in enumerate(KOLOM_EKSPOR, start=1):
        sheet.column_dimensions[get_column_letter(posisi)].width = lebar.get(label, 16)

    buffer = io.BytesIO()
    buku.save(buffer)
    return buffer.getvalue()
